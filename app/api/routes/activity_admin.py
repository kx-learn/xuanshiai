"""Activity and signup management for the independent back office (M7)."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from fastapi import APIRouter, Depends, HTTPException, Path, Query, Response
from openpyxl import Workbook
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.dependencies import CurrentMatchmakerAdmin, CurrentUser, get_current_admin, get_current_matchmaker_admin
from app.db.session import get_db
from app.schemas.activity_admin import (
    ActivityAdminCreate,
    ActivityAdminItem,
    ActivityAdminPage,
    ActivityAdminUpdate,
    ActivityLinkInfo,
    ActivityOption,
    ActivitySignupAdminItem,
    ActivitySignupAdminPage,
    ActivitySignupStatistics,
    ActivitySignupStatusUpdate,
    ActivitySignupUpdate,
    ActivityStatusUpdate,
)

router = APIRouter(prefix="/admin/activities")
signup_router = APIRouter(prefix="/admin/activity-signups")

_XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_ACTIVITY_BASE_URL = "https://www.xuanshi.com/subpages/active/index"

SELECT_ACTIVITY = "SELECT * FROM offline_activity"


def _bool(value: Any) -> bool:
    try:
        return bool(int(value))
    except (TypeError, ValueError):
        return bool(value)


def _num(value: Any) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _activity_item(row: Any) -> ActivityAdminItem:
    m = dict(row)
    return ActivityAdminItem(
        id=int(m["id"]),
        title=m["title"],
        cover=m.get("cover"),
        type=m.get("type"),
        city=m.get("city"),
        address=m.get("address"),
        start_time=m["start_time"],
        end_time=m["end_time"],
        signup_deadline=m.get("signup_deadline"),
        max_people=int(m.get("max_people") or 0),
        current_people=int(m.get("current_people") or 0),
        price=_num(m.get("price")),
        status=int(m.get("status") or 1),
        description=m.get("description"),
        created_by=m.get("created_by"),
        created_at=m["created_at"],
        organizer=m.get("organizer"),
        cover_small=m.get("cover_small"),
        fee_name=m.get("fee_name") or "报名费",
        price_male=_num(m.get("price_male")),
        price_female=_num(m.get("price_female")),
        signup_mode=m.get("signup_mode") or "anyone",
        require_realname=_bool(m.get("require_realname")),
        limit_mode=m.get("limit_mode") or "gender",
        max_male=int(m.get("max_male") or 0),
        max_female=int(m.get("max_female") or 0),
        virtual_people=int(m.get("virtual_people") or 0),
        virtual_female=int(m.get("virtual_female") or 0),
        hide_signup_count=_bool(m.get("hide_signup_count")),
        reward_promoter=_num(m.get("reward_promoter")),
        reward_service=_num(m.get("reward_service")),
        reward_partner=_num(m.get("reward_partner")),
        reminder_html=m.get("reminder_html"),
        service_wechat=m.get("service_wechat"),
        service_qr=m.get("service_qr"),
        virtual_views=int(m.get("virtual_views") or 0),
        sort_order=int(m.get("sort_order") or 0),
        custom_share=_bool(m.get("custom_share")),
        manager_ids=m.get("manager_ids"),
        notify_phones=m.get("notify_phones"),
        online=_bool(m.get("online") if m.get("online") is not None else 1),
        audit_status=m.get("audit_status") or "approved",
        male_count=int(m.get("male_count") or 0),
        female_count=int(m.get("female_count") or 0),
        link_url=f"{_ACTIVITY_BASE_URL}?id={int(m['id'])}",
    )


def _signup_item(row: Any) -> ActivitySignupAdminItem:
    m = dict(row)
    return ActivitySignupAdminItem(
        id=int(m["id"]),
        activity_id=int(m["activity_id"]),
        activity_title=m.get("activity_title"),
        user_id=int(m["user_id"]),
        nickname=m.get("nickname"),
        real_name=m.get("real_name"),
        phone=m.get("phone"),
        remark=m.get("remark"),
        status=int(m.get("status") or 0),
        cancel_reason=m.get("cancel_reason"),
        created_at=m["created_at"],
        updated_at=m["updated_at"],
        gender=m.get("gender"),
        age=m.get("age"),
        height=m.get("height"),
        education=m.get("education"),
        income=m.get("income"),
        marriage_status=m.get("marriage_status"),
        company=m.get("company"),
        avatar=m.get("avatar") or m.get("user_avatar"),
        id_card=m.get("id_card"),
        is_member=_bool(m.get("is_member")),
        is_realname=_bool(m.get("is_realname")),
        signup_times=int(m.get("signup_times") or 1),
        pay_status=m.get("pay_status") or "free",
        pay_amount=_num(m.get("pay_amount")),
        checked_in=_bool(m.get("checked_in")),
        in_crm=_bool(m.get("in_crm")),
        promoter_id=m.get("promoter_id"),
        promoter_name=m.get("promoter_name"),
    )


async def _get_activity(db: AsyncSession, activity_id: int) -> ActivityAdminItem:
    row = (
        await db.execute(
            text(
                "SELECT a.*, "
                "(SELECT COUNT(*) FROM activity_signup s JOIN users u ON u.id = s.user_id "
                " WHERE s.activity_id = a.id AND s.status = 1 "
                " AND COALESCE(s.gender, CASE u.gender WHEN 1 THEN '男' WHEN 2 THEN '女' END) = '男') AS male_count, "
                "(SELECT COUNT(*) FROM activity_signup s JOIN users u ON u.id = s.user_id "
                " WHERE s.activity_id = a.id AND s.status = 1 "
                " AND COALESCE(s.gender, CASE u.gender WHEN 1 THEN '男' WHEN 2 THEN '女' END) = '女') AS female_count "
                "FROM offline_activity a WHERE a.id = :id"
            ),
            {"id": activity_id},
        )
    ).mappings().first()
    if not row:
        raise HTTPException(404, detail="活动不存在")
    return _activity_item(row)


@router.get("", response_model=ActivityAdminPage, summary="查询活动列表")
async def activities(
    page: int = Query(1, ge=1, le=1000),
    page_size: int = Query(20, ge=1, le=100),
    status: int | None = Query(None, ge=1, le=5),
    city: str | None = Query(None, max_length=64),
    search: str | None = Query(None, max_length=128),
    online: bool | None = Query(None),
    audit_status: str | None = Query(None, pattern="^(pending|approved|rejected)$"),
    current: CurrentMatchmakerAdmin = Depends(get_current_matchmaker_admin),
    db: AsyncSession = Depends(get_db),
) -> ActivityAdminPage:
    current.require("community.activity.read")
    where = ["1=1"]
    params: dict = {"limit": page_size, "offset": (page - 1) * page_size}
    if status is not None:
        where.append("status = :status")
        params["status"] = status
    if city:
        where.append("city = :city")
        params["city"] = city
    if search:
        where.append("title LIKE CONCAT('%', :search, '%')")
        params["search"] = search
    if online is not None:
        where.append("COALESCE(online, 1) = :online")
        params["online"] = 1 if online else 0
    if audit_status:
        where.append("COALESCE(audit_status, 'approved') = :audit_status")
        params["audit_status"] = audit_status
    clause = " AND ".join(where)
    rows = await db.execute(
        text(
            "SELECT a.*, "
            "(SELECT COUNT(*) FROM activity_signup s JOIN users u ON u.id = s.user_id "
            " WHERE s.activity_id = a.id AND s.status = 1 "
            " AND COALESCE(s.gender, CASE u.gender WHEN 1 THEN '男' WHEN 2 THEN '女' END) = '男') AS male_count, "
            "(SELECT COUNT(*) FROM activity_signup s JOIN users u ON u.id = s.user_id "
            " WHERE s.activity_id = a.id AND s.status = 1 "
            " AND COALESCE(s.gender, CASE u.gender WHEN 1 THEN '男' WHEN 2 THEN '女' END) = '女') AS female_count "
            f"FROM offline_activity a WHERE {clause} ORDER BY a.id DESC LIMIT :limit OFFSET :offset"
        ),
        params,
    )
    count = await db.execute(
        text(f"SELECT COUNT(*) FROM offline_activity a WHERE {clause}"),
        {k: v for k, v in params.items() if k not in ("limit", "offset")},
    )
    total = int(count.scalar() or 0)
    return ActivityAdminPage(
        items=[_activity_item(row) for row in rows.mappings().all()],
        page=page,
        page_size=page_size,
        total=total,
        has_more=page * page_size < total,
    )


@router.post("", response_model=ActivityAdminItem, status_code=201, summary="创建活动")
async def create_activity(
    body: ActivityAdminCreate,
    current: CurrentMatchmakerAdmin = Depends(get_current_matchmaker_admin),
    db: AsyncSession = Depends(get_db),
) -> ActivityAdminItem:
    current.require("community.activity.manage")
    payload = body.model_dump()
    payload = {k: (int(v) if isinstance(v, bool) else v) for k, v in payload.items()}
    payload["created_by"] = current.account.id
    columns = ", ".join(payload.keys())
    placeholders = ", ".join(f":{key}" for key in payload)
    result = await db.execute(
        text(f"INSERT INTO offline_activity ({columns}) VALUES ({placeholders})"),
        payload,
    )
    await db.commit()
    return await _get_activity(db, int(result.lastrowid))


@router.get("/options", response_model=list[ActivityOption], summary="活动下拉字典")
async def activity_options(
    current: CurrentMatchmakerAdmin = Depends(get_current_matchmaker_admin),
    db: AsyncSession = Depends(get_db),
) -> list[ActivityOption]:
    current.require("community.activity.read")
    rows = await db.execute(text("SELECT id, title FROM offline_activity ORDER BY id DESC LIMIT 500"))
    return [ActivityOption(id=int(r[0]), title=str(r[1])) for r in rows.all()]


@router.get("/{activity_id}/signups", response_model=ActivitySignupAdminPage, summary="查询活动报名")
async def signups(
    activity_id: int = Path(..., ge=1),
    page: int = Query(1, ge=1, le=1000),
    page_size: int = Query(20, ge=1, le=100),
    status: int | None = Query(None, ge=0, le=3),
    current: CurrentMatchmakerAdmin = Depends(get_current_matchmaker_admin),
    db: AsyncSession = Depends(get_db),
) -> ActivitySignupAdminPage:
    current.require("community.activity.read")
    await _get_activity(db, activity_id)
    where = ["s.activity_id = :activity_id"]
    params: dict = {"activity_id": activity_id, "limit": page_size, "offset": (page - 1) * page_size}
    if status is not None:
        where.append("s.status = :status")
        params["status"] = status
    clause = " AND ".join(where)
    rows = await db.execute(
        text(f"SELECT s.*, u.nickname, u.avatar AS user_avatar, a.title AS activity_title, p.nickname AS promoter_name "
             f"FROM activity_signup s LEFT JOIN users u ON u.id = s.user_id "
             f"LEFT JOIN offline_activity a ON a.id = s.activity_id "
             f"LEFT JOIN users p ON p.id = s.promoter_id "
             f"WHERE {clause} ORDER BY s.id DESC LIMIT :limit OFFSET :offset"),
        params,
    )
    count = await db.execute(
        text(f"SELECT COUNT(*) FROM activity_signup s WHERE {clause}"),
        {k: v for k, v in params.items() if k not in ("limit", "offset")},
    )
    total = int(count.scalar() or 0)
    return ActivitySignupAdminPage(
        items=[_signup_item(row) for row in rows.mappings().all()],
        page=page,
        page_size=page_size,
        total=total,
        has_more=page * page_size < total,
    )


@router.get("/{activity_id}/link", response_model=ActivityLinkInfo, summary="活动链接与二维码")
async def activity_link(
    activity_id: int = Path(..., ge=1),
    current: CurrentMatchmakerAdmin = Depends(get_current_matchmaker_admin),
) -> ActivityLinkInfo:
    current.require("community.activity.read")
    return ActivityLinkInfo(link_url=f"{_ACTIVITY_BASE_URL}?id={activity_id}", qr_code=None)


@router.post("/{activity_id}/copy", response_model=ActivityAdminItem, status_code=201, summary="复制活动")
async def copy_activity(
    activity_id: int = Path(..., ge=1),
    current: CurrentMatchmakerAdmin = Depends(get_current_matchmaker_admin),
    db: AsyncSession = Depends(get_db),
) -> ActivityAdminItem:
    current.require("community.activity.manage")
    await _get_activity(db, activity_id)
    await db.execute(
        text(
            "INSERT INTO offline_activity "
            "(title, cover, cover_small, type, city, address, start_time, end_time, signup_deadline, "
            " max_people, price, description, organizer, time_text, fee_name, price_male, price_female, signup_mode, "
            " require_realname, limit_mode, max_male, max_female, virtual_people, virtual_female, "
            " hide_signup_count, reward_promoter, reward_service, reward_partner, reminder_html, "
            " service_wechat, service_qr, virtual_views, sort_order, custom_share, manager_ids, "
            " notify_phones, status, online, audit_status, created_by) "
            "SELECT CONCAT(title, '(副本)'), cover, cover_small, type, city, address, start_time, end_time, "
            " signup_deadline, max_people, price, description, organizer, time_text, fee_name, price_male, price_female, "
            " signup_mode, require_realname, limit_mode, max_male, max_female, virtual_people, virtual_female, "
            " hide_signup_count, reward_promoter, reward_service, reward_partner, reminder_html, "
            " service_wechat, service_qr, virtual_views, sort_order, custom_share, manager_ids, "
            " notify_phones, 1, 0, 'pending', :actor "
            "FROM offline_activity WHERE id = :id"
        ),
        {"id": activity_id, "actor": current.account.id},
    )
    await db.commit()
    fresh = (await db.execute(text("SELECT MAX(id) AS id FROM offline_activity"))).scalar()
    return await _get_activity(db, int(fresh))


@router.patch("/{activity_id}/status", response_model=ActivityAdminItem, summary="修改活动状态")
async def update_activity_status(
    activity_id: int = Path(..., ge=1),
    body: ActivityStatusUpdate = ...,
    current: CurrentMatchmakerAdmin = Depends(get_current_matchmaker_admin),
    db: AsyncSession = Depends(get_db),
) -> ActivityAdminItem:
    current.require("community.activity.manage")
    await _get_activity(db, activity_id)
    await db.execute(
        text("UPDATE offline_activity SET status = :status, updated_at = UTC_TIMESTAMP() WHERE id = :id"),
        {"status": body.status, "id": activity_id},
    )
    await db.commit()
    return await _get_activity(db, activity_id)


@router.patch("/{activity_id}", response_model=ActivityAdminItem, summary="修改活动")
async def update_activity(
    activity_id: int = Path(..., ge=1),
    body: ActivityAdminUpdate = ...,
    current: CurrentMatchmakerAdmin = Depends(get_current_matchmaker_admin),
    db: AsyncSession = Depends(get_db),
) -> ActivityAdminItem:
    current.require("community.activity.manage")
    await _get_activity(db, activity_id)
    values = {k: (int(v) if isinstance(v, bool) else v) for k, v in body.model_dump(exclude_unset=True).items()}
    updates = ", ".join(f"{key} = :{key}" for key in values)
    await db.execute(
        text(f"UPDATE offline_activity SET {updates}, updated_at = UTC_TIMESTAMP() WHERE id = :id"),
        {**values, "id": activity_id},
    )
    await db.commit()
    return await _get_activity(db, activity_id)


@router.delete("/{activity_id}", status_code=204, summary="删除活动")
async def delete_activity(
    activity_id: int = Path(..., ge=1),
    current: CurrentMatchmakerAdmin = Depends(get_current_matchmaker_admin),
    db: AsyncSession = Depends(get_db),
) -> None:
    current.require("community.activity.manage")
    await _get_activity(db, activity_id)
    await db.execute(text("DELETE FROM activity_signup WHERE activity_id = :id"), {"id": activity_id})
    await db.execute(text("DELETE FROM offline_activity WHERE id = :id"), {"id": activity_id})
    await db.execute(
        text("INSERT INTO business_audit_log (actor_user_id, action, resource_type, resource_id) "
             "VALUES (:actor, 'activity.delete', 'offline_activity', :id)"),
        {"actor": current.account.id, "id": activity_id},
    )
    await db.commit()


@router.get("/{activity_id}", response_model=ActivityAdminItem, summary="查询活动详情")
async def activity_detail(
    activity_id: int = Path(..., ge=1),
    current: CurrentMatchmakerAdmin = Depends(get_current_matchmaker_admin),
    db: AsyncSession = Depends(get_db),
) -> ActivityAdminItem:
    current.require("community.activity.read")
    return await _get_activity(db, activity_id)


# ─────────────────────────── 报名管理（独立前缀） ───────────────────────────


@signup_router.get("/statistics", response_model=ActivitySignupStatistics, summary="活动报名统计卡")
async def signup_statistics(
    activity_id: int | None = Query(None, ge=1),
    current: CurrentMatchmakerAdmin = Depends(get_current_matchmaker_admin),
    db: AsyncSession = Depends(get_db),
) -> ActivitySignupStatistics:
    current.require("community.activity.read")
    where = "WHERE 1=1"
    params: dict = {}
    if activity_id:
        where += " AND activity_id = :activity_id"
        params["activity_id"] = activity_id
    row = (
        await db.execute(
            text(
                "SELECT COUNT(*) AS total, "
                "SUM(CASE WHEN signup_times <= 1 THEN 1 ELSE 0 END) AS first_signup, "
                "SUM(CASE WHEN status = 0 THEN 1 ELSE 0 END) AS pending, "
                "SUM(CASE WHEN status = 1 THEN 1 ELSE 0 END) AS approved, "
                "SUM(CASE WHEN status = 3 THEN 1 ELSE 0 END) AS rejected, "
                "COALESCE(SUM(pay_amount), 0) AS fee_amount, "
                "SUM(CASE WHEN COALESCE(checked_in,0) = 0 THEN 1 ELSE 0 END) AS not_checked_in, "
                "SUM(CASE WHEN COALESCE(checked_in,0) = 1 THEN 1 ELSE 0 END) AS checked_in, "
                "SUM(CASE WHEN COALESCE(in_crm,0) = 1 THEN 1 ELSE 0 END) AS in_crm "
                f"FROM activity_signup {where}"
            ),
            params,
        )
    ).mappings().first()
    m = dict(row or {})
    return ActivitySignupStatistics(
        total=int(m.get("total") or 0),
        first_signup=int(m.get("first_signup") or 0),
        pending=int(m.get("pending") or 0),
        approved=int(m.get("approved") or 0),
        rejected=int(m.get("rejected") or 0),
        fee_amount=format(_num(m.get("fee_amount")), ".2f"),
        not_checked_in=int(m.get("not_checked_in") or 0),
        checked_in=int(m.get("checked_in") or 0),
        in_crm=int(m.get("in_crm") or 0),
    )


@signup_router.get("/options", response_model=list[ActivityOption], summary="报名筛选活动下拉")
async def signup_options(
    current: CurrentMatchmakerAdmin = Depends(get_current_matchmaker_admin),
    db: AsyncSession = Depends(get_db),
) -> list[ActivityOption]:
    current.require("community.activity.read")
    rows = await db.execute(text("SELECT id, title FROM offline_activity ORDER BY id DESC LIMIT 500"))
    return [ActivityOption(id=int(r[0]), title=str(r[1])) for r in rows.all()]


@signup_router.get("/export", summary="导出活动报名 Excel")
async def signup_export(
    activity_id: int | None = Query(None, ge=1),
    status: int | None = Query(None, ge=0, le=3),
    current: CurrentMatchmakerAdmin = Depends(get_current_matchmaker_admin),
    db: AsyncSession = Depends(get_db),
) -> Response:
    current.require("community.activity.read")
    where = ["1=1"]
    params: dict = {}
    if activity_id:
        where.append("s.activity_id = :activity_id")
        params["activity_id"] = activity_id
    if status is not None:
        where.append("s.status = :status")
        params["status"] = status
    clause = " AND ".join(where)
    rows = (
        await db.execute(
            text(
                "SELECT s.id, a.title AS activity_title, s.user_id, u.nickname, s.real_name, s.gender, "
                "s.age, s.phone, s.signup_times, s.pay_status, s.pay_amount, s.status, s.created_at "
                "FROM activity_signup s LEFT JOIN users u ON u.id = s.user_id "
                "LEFT JOIN offline_activity a ON a.id = s.activity_id "
                f"WHERE {clause} ORDER BY s.id DESC"
            ),
            params,
        )
    ).mappings().all()
    status_label = {0: "待审", 1: "审核通过", 2: "已取消", 3: "未通过"}
    wb = Workbook()
    ws = wb.active
    ws.title = "活动报名"
    ws.append(["ID", "活动", "会员ID", "昵称", "姓名", "性别", "年龄", "手机", "第几次报名", "缴费状态", "缴费金额", "审核状态", "报名时间"])
    for r in rows:
        ws.append([
            int(r["id"]), r["activity_title"], int(r["user_id"]), r["nickname"], r["real_name"],
            r["gender"], r["age"], r["phone"], int(r["signup_times"] or 1), r["pay_status"],
            _num(r["pay_amount"]), status_label.get(int(r["status"] or 0), "待审"),
            r["created_at"].strftime("%Y-%m-%d %H:%M:%S") if r["created_at"] else "",
        ])
    buffer = BytesIO()
    wb.save(buffer)
    return Response(
        content=buffer.getvalue(),
        media_type=_XLSX_MEDIA_TYPE,
        headers={"Content-Disposition": 'attachment; filename="activity-signups.xlsx"'},
    )


@signup_router.get("/{signup_id}", response_model=ActivitySignupAdminItem, summary="查询活动报名详情")
async def signup_detail(
    signup_id: int = Path(..., ge=1),
    current: CurrentMatchmakerAdmin = Depends(get_current_matchmaker_admin),
    db: AsyncSession = Depends(get_db),
) -> ActivitySignupAdminItem:
    current.require("community.activity.read")
    row = (
        await db.execute(
            text("SELECT s.*, u.nickname, u.avatar AS user_avatar, a.title AS activity_title, p.nickname AS promoter_name "
                 "FROM activity_signup s LEFT JOIN users u ON u.id = s.user_id "
                 "LEFT JOIN offline_activity a ON a.id = s.activity_id "
                 "LEFT JOIN users p ON p.id = s.promoter_id WHERE s.id = :id"),
            {"id": signup_id},
        )
    ).mappings().first()
    if not row:
        raise HTTPException(404, detail="活动报名不存在")
    return _signup_item(row)


@signup_router.patch("/{signup_id}", response_model=ActivitySignupAdminItem, summary="修改报名/审核报名")
async def update_signup(
    signup_id: int = Path(..., ge=1),
    body: ActivitySignupUpdate = ...,
    current: CurrentMatchmakerAdmin = Depends(get_current_matchmaker_admin),
    db: AsyncSession = Depends(get_db),
) -> ActivitySignupAdminItem:
    current.require("community.activity.manage")
    values = {k: (int(v) if isinstance(v, bool) else v) for k, v in body.model_dump(exclude_unset=True).items()}
    if not values:
        raise HTTPException(422, detail="至少提供一个需要修改的字段")
    updates = ", ".join(f"{key} = :{key}" for key in values)
    result = await db.execute(
        text(f"UPDATE activity_signup SET {updates}, updated_at = UTC_TIMESTAMP() WHERE id = :id"),
        {**values, "id": signup_id},
    )
    if result.rowcount == 0:
        raise HTTPException(404, detail="活动报名不存在")
    await db.execute(
        text("INSERT INTO business_audit_log (actor_user_id, action, resource_type, resource_id) "
             "VALUES (:actor, 'activity_signup.update', 'activity_signup', :id)"),
        {"actor": current.account.id, "id": signup_id},
    )
    await db.commit()
    row = (
        await db.execute(
            text("SELECT s.*, u.nickname, u.avatar AS user_avatar, a.title AS activity_title, p.nickname AS promoter_name "
                 "FROM activity_signup s LEFT JOIN users u ON u.id = s.user_id "
                 "LEFT JOIN offline_activity a ON a.id = s.activity_id "
                 "LEFT JOIN users p ON p.id = s.promoter_id WHERE s.id = :id"),
            {"id": signup_id},
        )
    ).mappings().one()
    return _signup_item(row)


@signup_router.delete("/{signup_id}", status_code=204, summary="删除活动报名")
async def delete_signup(
    signup_id: int = Path(..., ge=1),
    current: CurrentMatchmakerAdmin = Depends(get_current_matchmaker_admin),
    db: AsyncSession = Depends(get_db),
) -> None:
    current.require("community.activity.manage")
    result = await db.execute(text("DELETE FROM activity_signup WHERE id = :id"), {"id": signup_id})
    if result.rowcount == 0:
        raise HTTPException(404, detail="活动报名不存在")
    await db.commit()
