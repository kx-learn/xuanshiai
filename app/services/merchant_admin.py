"""Merchant alliance (商家联盟) service for the back office (M7-B)."""

from __future__ import annotations

import json
import uuid
from datetime import datetime, timezone
from typing import Any

from fastapi import HTTPException
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.schemas.merchant_admin import (
    MerchantCategoryCreate,
    MerchantCategoryItem,
    MerchantCategoryUpdate,
    MerchantCreate,
    MerchantItem,
    MerchantOrderItem,
    MerchantOrderPage,
    MerchantOrderStatusUpdate,
    MerchantPage,
    MerchantProductCreate,
    MerchantProductItem,
    MerchantProductPage,
    MerchantProductUpdate,
    MerchantUpdate,
)

_BASE_URL = "https://www.xuanshi.com/subpages/hezuo"


def _num(value: Any) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _money(value: Any) -> str:
    return format(_num(value), ".2f")


def _bool(value: Any) -> bool:
    try:
        return bool(int(value))
    except (TypeError, ValueError):
        return bool(value)


def _gallery(value: Any) -> list[str]:
    if not value:
        return []
    try:
        data = json.loads(value) if isinstance(value, str) else value
        return [str(x) for x in data] if isinstance(data, list) else []
    except (TypeError, ValueError):
        return []


def _tags(value: Any) -> list[str]:
    if not value:
        return []
    return [t for t in str(value).split(",") if t]


# ─────────────────────────── 分类 ───────────────────────────


def _cat_item(row: Any) -> MerchantCategoryItem:
    m = dict(row)
    return MerchantCategoryItem(
        id=int(m["id"]),
        name=m["name"],
        icon_url=m.get("icon_url"),
        sort=int(m.get("sort") or 0),
        status=int(m.get("status") or 0),
        merchant_count=int(m.get("merchant_count") or 0),
    )


async def list_categories(db: AsyncSession) -> list[MerchantCategoryItem]:
    rows = await db.execute(
        text(
            "SELECT c.*, (SELECT COUNT(*) FROM merchant m WHERE m.category_id = c.id AND m.deleted_at IS NULL) "
            "AS merchant_count FROM merchant_category c ORDER BY c.sort ASC, c.id ASC"
        )
    )
    return [_cat_item(r) for r in rows.mappings().all()]


async def create_category(db: AsyncSession, body: MerchantCategoryCreate) -> MerchantCategoryItem:
    exists = (await db.execute(text("SELECT id FROM merchant_category WHERE name = :n"), {"n": body.name})).scalar()
    if exists:
        raise HTTPException(409, detail="分类名称已存在")
    result = await db.execute(
        text("INSERT INTO merchant_category (name, icon_url, sort, status) VALUES (:n, :i, :s, :st)"),
        {"n": body.name, "i": body.icon_url, "s": body.sort, "st": body.status},
    )
    await db.commit()
    return await _get_category(db, int(result.lastrowid))


async def _get_category(db: AsyncSession, category_id: int) -> MerchantCategoryItem:
    row = (
        await db.execute(
            text("SELECT c.*, 0 AS merchant_count FROM merchant_category c WHERE c.id = :id"), {"id": category_id}
        )
    ).mappings().first()
    if not row:
        raise HTTPException(404, detail="商家分类不存在")
    return _cat_item(row)


async def update_category(db: AsyncSession, category_id: int, body: MerchantCategoryUpdate) -> MerchantCategoryItem:
    await _get_category(db, category_id)
    values = body.model_dump(exclude_unset=True)
    if "name" in values:
        dup = (
            await db.execute(
                text("SELECT id FROM merchant_category WHERE name = :n AND id <> :id"),
                {"n": values["name"], "id": category_id},
            )
        ).scalar()
        if dup:
            raise HTTPException(409, detail="分类名称已存在")
    updates = ", ".join(f"{k} = :{k}" for k in values)
    await db.execute(
        text(f"UPDATE merchant_category SET {updates}, updated_at = UTC_TIMESTAMP() WHERE id = :id"),
        {**values, "id": category_id},
    )
    await db.commit()
    return await _get_category(db, category_id)


async def delete_category(db: AsyncSession, category_id: int) -> None:
    await _get_category(db, category_id)
    used = (
        await db.execute(
            text("SELECT COUNT(*) FROM merchant WHERE category_id = :id AND deleted_at IS NULL"), {"id": category_id}
        )
    ).scalar()
    if int(used or 0) > 0:
        raise HTTPException(409, detail="该分类下仍有商家，无法删除")
    await db.execute(text("DELETE FROM merchant_category WHERE id = :id"), {"id": category_id})
    await db.commit()


async def reorder_categories(db: AsyncSession, ids: list[int]) -> None:
    for index, cid in enumerate(ids, start=1):
        await db.execute(
            text("UPDATE merchant_category SET sort = :s WHERE id = :id"), {"s": index, "id": cid}
        )
    await db.commit()


# ─────────────────────────── 商家 ───────────────────────────

_MERCHANT_SELECT = """
SELECT m.*, c.name AS category_name, u.nickname AS admin_user_nickname,
    (SELECT COUNT(*) FROM merchant_product p WHERE p.merchant_id = m.id AND p.deleted_at IS NULL) AS product_total,
    (SELECT COUNT(*) FROM merchant_product p WHERE p.merchant_id = m.id AND p.deleted_at IS NULL AND p.status = 1) AS product_online,
    (SELECT COALESCE(SUM(o.amount), 0) FROM merchant_order o
      WHERE o.merchant_id = m.id AND o.verify_status = 'verified') AS sales_amount
FROM merchant m
LEFT JOIN merchant_category c ON c.id = m.category_id
LEFT JOIN users u ON u.id = m.admin_user_id
"""


def _merchant_item(row: Any) -> MerchantItem:
    m = dict(row)
    return MerchantItem(
        id=int(m["id"]),
        name=m["name"],
        cover=m.get("cover"),
        gallery=_gallery(m.get("gallery_json")),
        category_id=m.get("category_id"),
        category_name=m.get("category_name"),
        tags=_tags(m.get("tags")),
        province=m.get("province"),
        city=m.get("city"),
        address=m.get("address"),
        contact_phone=m.get("contact_phone"),
        business_hours=m.get("business_hours"),
        intro=m.get("intro"),
        admin_user_id=m.get("admin_user_id"),
        admin_user_nickname=m.get("admin_user_nickname"),
        verify_staff=m.get("admin_user_nickname"),
        sort=int(m.get("sort") or 0),
        visible=_bool(m.get("visible") if m.get("visible") is not None else 1),
        product_total=int(m.get("product_total") or 0),
        product_online=int(m.get("product_online") or 0),
        sales_amount=_money(m.get("sales_amount")),
        created_at=m.get("created_at"),
        link_url=f"{_BASE_URL}/merchant?id={int(m['id'])}",
        qr_code=m.get("qr_code"),
    )


async def _get_merchant(db: AsyncSession, merchant_id: int) -> MerchantItem:
    row = (
        await db.execute(text(f"{_MERCHANT_SELECT} WHERE m.id = :id AND m.deleted_at IS NULL"), {"id": merchant_id})
    ).mappings().first()
    if not row:
        raise HTTPException(404, detail="商家不存在")
    return _merchant_item(row)


def _merchant_payload(body: MerchantCreate | MerchantUpdate, partial: bool) -> dict[str, Any]:
    data = body.model_dump(exclude_unset=partial)
    out: dict[str, Any] = {}
    for key, value in data.items():
        if key == "gallery":
            out["gallery_json"] = json.dumps(value or [], ensure_ascii=False)
        elif key == "tags":
            out["tags"] = ",".join(value or [])
        elif key == "visible":
            out["visible"] = 1 if value else 0
        else:
            out[key] = value
    return out


async def list_merchants(
    db: AsyncSession, page: int, page_size: int, category_id: int | None, keyword: str | None, visible: bool | None
) -> MerchantPage:
    where = ["m.deleted_at IS NULL"]
    params: dict = {"limit": page_size, "offset": (page - 1) * page_size}
    if category_id:
        where.append("m.category_id = :category_id")
        params["category_id"] = category_id
    if keyword:
        where.append("m.name LIKE CONCAT('%', :kw, '%')")
        params["kw"] = keyword
    if visible is not None:
        where.append("m.visible = :visible")
        params["visible"] = 1 if visible else 0
    clause = " AND ".join(where)
    rows = await db.execute(
        text(f"{_MERCHANT_SELECT} WHERE {clause} ORDER BY m.sort DESC, m.id DESC LIMIT :limit OFFSET :offset"), params
    )
    count = await db.execute(
        text(f"SELECT COUNT(*) FROM merchant m WHERE {clause}"),
        {k: v for k, v in params.items() if k not in ("limit", "offset")},
    )
    total = int(count.scalar() or 0)
    return MerchantPage(
        items=[_merchant_item(r) for r in rows.mappings().all()],
        page=page,
        page_size=page_size,
        total=total,
        has_more=page * page_size < total,
    )


async def create_merchant(db: AsyncSession, body: MerchantCreate) -> MerchantItem:
    payload = _merchant_payload(body, partial=False)
    payload["link_url"] = None
    columns = ", ".join(payload.keys())
    placeholders = ", ".join(f":{k}" for k in payload)
    result = await db.execute(text(f"INSERT INTO merchant ({columns}) VALUES ({placeholders})"), payload)
    await db.commit()
    return await _get_merchant(db, int(result.lastrowid))


async def update_merchant(db: AsyncSession, merchant_id: int, body: MerchantUpdate) -> MerchantItem:
    await _get_merchant(db, merchant_id)
    values = _merchant_payload(body, partial=True)
    updates = ", ".join(f"{k} = :{k}" for k in values)
    await db.execute(
        text(f"UPDATE merchant SET {updates}, updated_at = UTC_TIMESTAMP() WHERE id = :id"),
        {**values, "id": merchant_id},
    )
    await db.commit()
    return await _get_merchant(db, merchant_id)


async def set_merchant_visible(db: AsyncSession, merchant_id: int, visible: bool) -> MerchantItem:
    await _get_merchant(db, merchant_id)
    await db.execute(
        text("UPDATE merchant SET visible = :v, updated_at = UTC_TIMESTAMP() WHERE id = :id"),
        {"v": 1 if visible else 0, "id": merchant_id},
    )
    await db.commit()
    return await _get_merchant(db, merchant_id)


async def delete_merchant(db: AsyncSession, merchant_id: int) -> None:
    await _get_merchant(db, merchant_id)
    await db.execute(
        text("UPDATE merchant SET deleted_at = UTC_TIMESTAMP() WHERE id = :id"), {"id": merchant_id}
    )
    await db.commit()


# ─────────────────────────── 商品 ───────────────────────────

_PRODUCT_SELECT = """
SELECT p.*, m.name AS merchant_name,
    (SELECT COUNT(*) FROM merchant_order o WHERE o.product_id = p.id AND o.pay_status = 'paid') AS sales_count,
    (SELECT COALESCE(SUM(o.amount), 0) FROM merchant_order o
      WHERE o.product_id = p.id AND o.verify_status = 'verified') AS sales_amount
FROM merchant_product p
LEFT JOIN merchant m ON m.id = p.merchant_id
"""


def _product_item(row: Any) -> MerchantProductItem:
    m = dict(row)
    return MerchantProductItem(
        id=int(m["id"]),
        merchant_id=int(m["merchant_id"]),
        merchant_name=m.get("merchant_name"),
        name=m["name"],
        cover=m.get("cover"),
        original_price=_num(m.get("original_price")),
        sale_price=_num(m.get("sale_price")),
        settle_amount=_num(m.get("settle_amount")),
        promote_split_mode=m.get("promote_split_mode") or "fixed",
        promote_amount=_num(m.get("promote_amount")),
        partner_split_mode=m.get("partner_split_mode") or "fixed",
        partner_amount=_num(m.get("partner_amount")),
        service_amount=_num(m.get("service_amount")),
        buy_limit_mode=m.get("buy_limit_mode") or "account",
        account_limit=int(m.get("account_limit") or 0),
        order_limit=int(m.get("order_limit") or 0),
        notice_mode=m.get("notice_mode") or "default",
        notice_text=m.get("notice_text"),
        intro=m.get("intro"),
        status=int(m.get("status") or 1),
        sort=int(m.get("sort") or 0),
        sales_count=int(m.get("sales_count") or 0),
        sales_amount=_money(m.get("sales_amount")),
        create_time=m.get("created_at"),
        created_at=m.get("created_at"),
        link_url=f"{_BASE_URL}/product?id={int(m['id'])}",
        qr_code=m.get("qr_code"),
    )


async def _get_product(db: AsyncSession, product_id: int) -> MerchantProductItem:
    row = (
        await db.execute(text(f"{_PRODUCT_SELECT} WHERE p.id = :id AND p.deleted_at IS NULL"), {"id": product_id})
    ).mappings().first()
    if not row:
        raise HTTPException(404, detail="商品不存在")
    return _product_item(row)


async def list_products(
    db: AsyncSession, page: int, page_size: int, merchant_id: int | None, keyword: str | None, status: int | None
) -> MerchantProductPage:
    where = ["p.deleted_at IS NULL"]
    params: dict = {"limit": page_size, "offset": (page - 1) * page_size}
    if merchant_id:
        where.append("p.merchant_id = :merchant_id")
        params["merchant_id"] = merchant_id
    if keyword:
        where.append("p.name LIKE CONCAT('%', :kw, '%')")
        params["kw"] = keyword
    if status is not None:
        where.append("p.status = :status")
        params["status"] = status
    clause = " AND ".join(where)
    rows = await db.execute(
        text(f"{_PRODUCT_SELECT} WHERE {clause} ORDER BY p.id DESC LIMIT :limit OFFSET :offset"), params
    )
    count = await db.execute(
        text(f"SELECT COUNT(*) FROM merchant_product p WHERE {clause}"),
        {k: v for k, v in params.items() if k not in ("limit", "offset")},
    )
    total = int(count.scalar() or 0)
    return MerchantProductPage(
        items=[_product_item(r) for r in rows.mappings().all()],
        page=page,
        page_size=page_size,
        total=total,
        has_more=page * page_size < total,
    )


async def create_product(db: AsyncSession, body: MerchantProductCreate) -> MerchantProductItem:
    await _get_merchant(db, body.merchant_id)
    payload = body.model_dump()
    columns = ", ".join(payload.keys())
    placeholders = ", ".join(f":{k}" for k in payload)
    result = await db.execute(
        text(f"INSERT INTO merchant_product ({columns}) VALUES ({placeholders})"), payload
    )
    await db.commit()
    return await _get_product(db, int(result.lastrowid))


async def update_product(db: AsyncSession, product_id: int, body: MerchantProductUpdate) -> MerchantProductItem:
    await _get_product(db, product_id)
    values = body.model_dump(exclude_unset=True)
    if "merchant_id" in values:
        await _get_merchant(db, values["merchant_id"])
    updates = ", ".join(f"{k} = :{k}" for k in values)
    await db.execute(
        text(f"UPDATE merchant_product SET {updates}, updated_at = UTC_TIMESTAMP() WHERE id = :id"),
        {**values, "id": product_id},
    )
    await db.commit()
    return await _get_product(db, product_id)


async def set_product_status(db: AsyncSession, product_id: int, status: int) -> MerchantProductItem:
    await _get_product(db, product_id)
    await db.execute(
        text("UPDATE merchant_product SET status = :s, updated_at = UTC_TIMESTAMP() WHERE id = :id"),
        {"s": status, "id": product_id},
    )
    await db.commit()
    return await _get_product(db, product_id)


async def delete_product(db: AsyncSession, product_id: int) -> None:
    await _get_product(db, product_id)
    await db.execute(
        text("UPDATE merchant_product SET deleted_at = UTC_TIMESTAMP() WHERE id = :id"), {"id": product_id}
    )
    await db.commit()


# ─────────────────────────── 订单 ───────────────────────────


def _order_item(row: Any) -> MerchantOrderItem:
    m = dict(row)
    return MerchantOrderItem(
        id=int(m["id"]),
        order_no=m["order_no"],
        product_id=int(m["product_id"]),
        product_name=m.get("product_name"),
        product_cover=m.get("product_cover"),
        merchant_id=int(m["merchant_id"]),
        merchant_name=m.get("merchant_name"),
        buyer_user_id=int(m["buyer_user_id"]),
        buyer_nickname=m.get("buyer_nickname"),
        buyer_phone=m.get("buyer_phone"),
        quantity=int(m.get("quantity") or 1),
        amount=_money(m.get("amount")),
        pay_status=m.get("pay_status") or "unpaid",
        pay_method=m.get("pay_method"),
        paid_at=m.get("paid_at"),
        verify_status=m.get("verify_status") or "pending",
        verify_code=m.get("verify_code"),
        verified_at=m.get("verified_at"),
        status=m.get("status") or "pending",
        remark=m.get("remark"),
        created_at=m.get("created_at"),
    )


_ORDER_SELECT = """
SELECT o.*, p.name AS product_name, p.cover AS product_cover, m.name AS merchant_name,
       u.nickname AS buyer_nickname, u.phone AS buyer_phone
FROM merchant_order o
LEFT JOIN merchant_product p ON p.id = o.product_id
LEFT JOIN merchant m ON m.id = o.merchant_id
LEFT JOIN users u ON u.id = o.buyer_user_id
"""


async def list_orders(
    db: AsyncSession,
    page: int,
    page_size: int,
    merchant_id: int | None,
    status: str | None,
    verify_status: str | None,
    keyword: str | None,
    order_no: str | None,
    buyer_keyword: str | None,
    verify_start: str | None,
    verify_end: str | None,
) -> MerchantOrderPage:
    where = ["1=1"]
    params: dict = {"limit": page_size, "offset": (page - 1) * page_size}
    if merchant_id:
        where.append("o.merchant_id = :merchant_id")
        params["merchant_id"] = merchant_id
    if status:
        where.append("o.status = :status")
        params["status"] = status
    if verify_status:
        where.append("o.verify_status = :verify_status")
        params["verify_status"] = verify_status
    if keyword:
        where.append("p.name LIKE CONCAT('%', :kw, '%')")
        params["kw"] = keyword
    if order_no:
        where.append("o.order_no LIKE CONCAT('%', :order_no, '%')")
        params["order_no"] = order_no
    if buyer_keyword:
        where.append("(u.nickname LIKE CONCAT('%', :bk, '%') OR u.phone LIKE CONCAT('%', :bk, '%'))")
        params["bk"] = buyer_keyword
    if verify_start:
        where.append("o.verified_at >= :verify_start")
        params["verify_start"] = verify_start
    if verify_end:
        where.append("o.verified_at < DATE_ADD(:verify_end, INTERVAL 1 DAY)")
        params["verify_end"] = verify_end
    clause = " AND ".join(where)
    rows = await db.execute(
        text(f"{_ORDER_SELECT} WHERE {clause} ORDER BY o.id DESC LIMIT :limit OFFSET :offset"), params
    )
    count = await db.execute(
        text(f"SELECT COUNT(*) FROM merchant_order o "
             f"LEFT JOIN merchant_product p ON p.id = o.product_id "
             f"LEFT JOIN users u ON u.id = o.buyer_user_id WHERE {clause}"),
        {k: v for k, v in params.items() if k not in ("limit", "offset")},
    )
    total = int(count.scalar() or 0)
    return MerchantOrderPage(
        items=[_order_item(r) for r in rows.mappings().all()],
        page=page,
        page_size=page_size,
        total=total,
        has_more=page * page_size < total,
    )


async def update_order_status(
    db: AsyncSession, order_id: int, body: MerchantOrderStatusUpdate
) -> MerchantOrderItem:
    row = (await db.execute(text("SELECT * FROM merchant_order WHERE id = :id"), {"id": order_id})).mappings().first()
    if not row:
        raise HTTPException(404, detail="订单不存在")
    current = str(row["status"])
    if current in ("used", "cancelled"):
        raise HTTPException(409, detail="已消费或已取消的订单不可变更")
    if body.status == "cancelled" and str(row["verify_status"]) == "verified":
        raise HTTPException(409, detail="已核销订单不可取消，请在财务管理中退款")
    await db.execute(
        text("UPDATE merchant_order SET status = :s, remark = :r, updated_at = UTC_TIMESTAMP() WHERE id = :id"),
        {"s": body.status, "r": body.remark, "id": order_id},
    )
    await db.commit()
    fresh = (await db.execute(text(f"{_ORDER_SELECT} WHERE o.id = :id"), {"id": order_id})).mappings().one()
    return _order_item(fresh)


def new_order_no() -> str:
    return f"MO{datetime.now(timezone.utc).strftime('%Y%m%d%H%M%S')}{uuid.uuid4().hex[:6].upper()}"


async def build_order_export(db: AsyncSession, **filters: Any) -> bytes:
    from io import BytesIO

    from openpyxl import Workbook

    result = await list_orders(db, page=1, page_size=10000, **filters)
    status_label = {"pending": "未支付", "paid": "已支付", "used": "已消费", "cancelled": "已取消"}
    wb = Workbook()
    ws = wb.active
    ws.title = "商家订单"
    ws.append(["ID", "订单号", "下单商品", "商家", "下单人", "下单时间", "订单状态", "订单金额", "支付方式", "支付时间", "核销状态", "核销时间"])
    for o in result.items:
        ws.append([
            o.id, o.order_no, o.product_name, o.merchant_name, o.buyer_nickname,
            o.created_at.strftime("%Y-%m-%d %H:%M:%S") if o.created_at else "",
            status_label.get(o.status, o.status), o.amount, o.pay_method,
            o.paid_at.strftime("%Y-%m-%d %H:%M:%S") if o.paid_at else "",
            "已核销" if o.verify_status == "verified" else "未核销",
            o.verified_at.strftime("%Y-%m-%d %H:%M:%S") if o.verified_at else "",
        ])
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
