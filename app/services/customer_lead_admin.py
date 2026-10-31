"""Customer lead management for the independent back office."""

import asyncio
import json
import re
import uuid
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from fastapi import HTTPException, UploadFile
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.sensitive_fields import REDACTION_LEVEL_STANDARD, is_masked_value, mask_contact
# 复用登录注册链路的默认角色发放实现，避免入库账号缺少默认 user_role。
from app.services.auth import ensure_default_user_role
from app.schemas.auth import PreferenceUpdateRequest

from app.schemas.customer_lead_admin import (
    CustomerLead, CustomerLeadAbandonment, CustomerLeadAssignment, CustomerLeadCreate, CustomerLeadCreateResult, CustomerLeadFollowUp,
    CustomerLeadBatchImportResult, CustomerLeadConvertRequest, CustomerLeadFollowUpCreate, CustomerLeadImportRow, CustomerLeadImportSummary,
    CustomerLeadOption, CustomerLeadOptions,
    CustomerLeadPage, CustomerLeadStatistics, CustomerLeadUpdate,
    CustomerLeadFollowUpRecord, CustomerLeadFollowUpRecordPage, CustomerLeadFollowBucket,
    MatchmakerFollowStat, CustomerLeadPromoterStat, LeadPreferenceResponse,
)
from app.services.customer_lead_contact import (
    ensure_contact_available, normalize_contact, raise_duplicate_contact,
)
from app.services.lead_apportion import resolve_assigned_matchmaker
from app.core.config import settings
from app.services.audio_duration import parse_audio_duration
from app.services.media_access import MEDIA_CATEGORY_FOLLOW_UP, sign_media_url
# 客源跟进附件与会员侧完全同构：直接复用会员跟进的落盘/webp/录音校验/JSON 解析助手，
# 避免两套实现漂移（图片转 webp、录音白名单、限流读取等均单一真相源）。
from app.services.member_follow_up_admin import (
    FOLLOW_UP_IMAGE_MAX_BYTES,
    FOLLOW_UP_MAX_IMAGES,
    FOLLOW_UP_VOICE_MAX_BYTES,
    _fu_image_outputs,
    _fu_read_limited,
    _fu_write_bytes,
    _parse_follow_up_images,
    _voice_suffix,
)

LEAD_SELECT = """SELECT id, name, phone, wechat, source, intention_level, status, audit_status, matchmaker_id,
    organization_id, promoter_id, next_follow_at, converted_user_id, remark, created_by, created_at, updated_at
    FROM customer_lead"""

_LEAD_FROM_RE = re.compile(r"\bFROM customer_lead\b")

# 客源「本人资料」可读写列（与 _LeadProfileFields 一一对应）。gender/age 单列处理：
# gender 沿用 CAST，age 由 birth_date 派生。列不存在时读写均静默降级（先发版后迁移）。
_LEAD_PROFILE_COLUMNS: tuple[str, ...] = (
    "real_name", "birth_date", "marriage", "occupation", "income", "height",
    "weight", "education", "hometown", "residence", "house", "car",
    "smoking", "drinking", "ethnicity", "school", "company", "photos",
)


def _lead_select_for(columns: frozenset[str]) -> str:
    """在客源主表外层 SELECT 补 gender 与已存在的本人资料列。

    age 优先由 birth_date 派生（TIMESTAMPDIFF），无 birth_date 时回退到 filters 迁移
    补的 raw age 列；两者都缺则不输出 age。缺列静默跳过，保证「先发版后迁移」
    不报 Unknown column。只在外层 ``FROM customer_lead`` 词边界处注入一次。
    """
    extra: list[str] = []
    if "gender" in columns:
        extra.append("CAST(gender AS SIGNED) AS gender")
    extra.extend(col for col in _LEAD_PROFILE_COLUMNS if col in columns)
    if "birth_date" in columns:
        extra.append("TIMESTAMPDIFF(YEAR, birth_date, UTC_DATE()) AS age")
    elif "age" in columns:
        extra.append("age")
    if not extra:
        return LEAD_SELECT
    return _LEAD_FROM_RE.sub(", " + ", ".join(extra) + " FROM customer_lead", LEAD_SELECT)


# 导入模板表头（顺序即模板列顺序）
_TEMPLATE_HEADERS: list[str] = ["姓名", "手机号", "微信号", "来源", "意向程度", "备注"]
_TEMPLATE_SAMPLE: list[str] = ["示例：张女士", "13800000000", "zhang_wx", "抖音", "B", "示例行，导入前请删除"]
# 客户意向七档（A-G），与后台筛选 / 首页意向统计标签同口径。
LEAD_INTENTION_GRADES: dict[int, str] = {
    1: "A类未接",
    2: "B类初步沟通",
    3: "C类深入沟通未缔结",
    4: "D类待确定到店时间",
    5: "E类已确定到店",
    6: "F类预约需二邀",
    7: "G类已到店未签约",
}
# 导入别名：数字 1-7 / 字母 A-G（大小写）/ 全称 / 旧粗粒度（低→A、中→D、高→F）。
_INTENTION_ALIASES: dict[str, int] = {
    **{str(g): g for g in LEAD_INTENTION_GRADES},
    **{chr(64 + g): g for g in LEAD_INTENTION_GRADES},
    **{chr(64 + g).lower(): g for g in LEAD_INTENTION_GRADES},
    **{label: g for g, label in LEAD_INTENTION_GRADES.items()},
    "低": 1,
    "中": 4,
    "高": 6,
}


def _lead(row: Any, redaction_level: str = REDACTION_LEVEL_STANDARD) -> CustomerLead:
    """客源线索序列化。

    `wechat` 按敏感等级掩码后返回（审计第 1 项：完整微信号不得返回）；默认标准级，
    调用方忘记传等级时落到最严格一档，不会泄露完整值。`photos` 为 JSON 列，
    raw SQL 取回的是文本，归一化为 URL 列表（NULL 保持 None）。
    """
    from app.services.profile import _json_list  # 延迟导入避免模块级循环依赖

    payload = dict(row)
    payload["wechat"] = mask_contact(payload.get("wechat"), redaction_level)
    if payload.get("photos") is not None:
        payload["photos"] = _json_list(payload["photos"])
    return CustomerLead(**payload)


async def _validate_owner(db: AsyncSession, assignment: CustomerLeadAssignment) -> None:
    if assignment.matchmaker_id is not None:
        result = await db.execute(text("""SELECT 1 FROM user_matchmaker_apply a JOIN user_role r ON r.user_id = a.user_id
            AND r.role_code = 'service_matchmaker' AND r.status = 1
            WHERE a.user_id = :id AND a.application_type = 'service_matchmaker' AND a.status = 1"""), {"id": assignment.matchmaker_id})
        if not result.scalar():
            raise HTTPException(422, detail="只能分配给有效服务红娘")
    if assignment.organization_id is not None:
        result = await db.execute(text("SELECT 1 FROM organization WHERE id = :id AND org_type = 'store' AND status = 1"), {"id": assignment.organization_id})
        if not result.scalar():
            raise HTTPException(422, detail="门店不存在或已停用")


async def _ensure_user_in_scope(db: AsyncSession, user_id: int | None, scope: str | None, scope_params: dict[str, Any] | None) -> None:
    if user_id is None:
        return
    if not scope:
        raise HTTPException(404, detail="目标人员不在当前数据范围")
    params = dict(scope_params or {})
    params["target_user_id"] = user_id
    if not await db.scalar(text(f"SELECT 1 FROM users WHERE id = :target_user_id AND ({scope})"), params):
        raise HTTPException(404, detail="目标人员不在当前数据范围")


async def create_lead(db: AsyncSession, account_id: int, request: CustomerLeadCreate, redaction_level: str = REDACTION_LEVEL_STANDARD, *, assignment_scope: str | None = None, assignment_scope_params: dict[str, Any] | None = None) -> CustomerLeadCreateResult:
    phone = normalize_contact(request.phone)
    wechat = normalize_contact(request.wechat)
    await _ensure_user_in_scope(db, request.promoter_id, assignment_scope, assignment_scope_params)
    await ensure_contact_available(db, phone, wechat)
    # auto_assign=true 时按分派配置在同一事务内选好归属红娘；选不到人（返回 None）则不写
    # matchmaker_id，客源保持未分派。策略计算全部在服务端完成。
    assigned_matchmaker_id = await resolve_assigned_matchmaker(db, auto_assign=request.auto_assign)
    from app.services.profile import _calculate_age, _json_value  # 延迟导入避免循环依赖
    try:
        columns = await _lead_columns(db)
        insert_columns = ["name", "phone", "wechat", "source", "intention_level", "remark", "promoter_id", "audit_status", "created_by"]
        insert_payload = {**request.model_dump(exclude={"auto_assign"}), "phone": phone, "wechat": wechat, "created_by": account_id}
        # photos 存 JSON 字符串；birth_date 有值时同步派生 raw age 列（供列表筛选与旧口径使用）。
        if insert_payload.get("photos") is not None:
            insert_payload["photos"] = _json_value(insert_payload["photos"])
        if insert_payload.get("birth_date") is not None:
            insert_payload["age"] = _calculate_age(insert_payload["birth_date"])
        # 基础列恒写；gender / 资料列 / age 仅写库中已存在且有值的列（缺列静默降级）。
        for col in ("gender", *_LEAD_PROFILE_COLUMNS, "age"):
            if col in columns and insert_payload.get(col) is not None:
                insert_columns.append(col)
        if assigned_matchmaker_id:
            insert_columns.append("matchmaker_id")
            insert_payload["matchmaker_id"] = assigned_matchmaker_id
            # 分派时间与归属同写：后台「分派日期」筛选只认 assigned_at，缺列时静默跳过。
            if "assigned_at" in columns:
                insert_columns.append("assigned_at")
                insert_payload["assigned_at"] = datetime.utcnow()
        insert_sql = f"""INSERT INTO customer_lead
            ({', '.join(insert_columns)})
            VALUES (:{', '.join(insert_columns)})"""
        result = await db.execute(text(insert_sql), insert_payload)
        lead_id = int(result.lastrowid)
        await db.execute(text("""INSERT INTO business_audit_log (actor_user_id, action, resource_type, resource_id)
            VALUES (:actor, 'customer_lead.create', 'customer_lead', :id)"""), {"actor": account_id, "id": lead_id})
        await db.commit()
    except IntegrityError:
        await db.rollback()
        raise_duplicate_contact()
    lead = await get_lead(db, lead_id, redaction_level)
    return CustomerLeadCreateResult(**lead.model_dump(), assigned=bool(assigned_matchmaker_id))


def _merge_scope(
    where: list[str], params: dict[str, Any], scope: str | None, scope_params: dict[str, Any] | None
) -> None:
    """把四档数据范围谓词并入查询（D-16）。

    ``scope`` 由 `CurrentMatchmakerAdmin.scope_clause` 在路由层生成（唯一实现），
    服务层只负责拼接；``scope`` 为空表示不做约束（超管 / 服务层直调）。
    """
    if not scope:
        return
    params.update(scope_params or {})
    where.append(f"({scope})")


async def get_lead(
    db: AsyncSession,
    lead_id: int,
    redaction_level: str = REDACTION_LEVEL_STANDARD,
    *,
    scope: str | None = None,
    scope_params: dict[str, Any] | None = None,
) -> CustomerLead:
    where = ["id = :id"]
    params: dict[str, Any] = {"id": lead_id}
    _merge_scope(where, params, scope, scope_params)
    columns = await _lead_columns(db)
    if "deleted_at" in columns:
        where.append("deleted_at IS NULL")
    select_sql = _lead_select_for(columns)
    row = (await db.execute(text(f"{select_sql} WHERE {' AND '.join(where)}"), params)).mappings().first()
    if not row:
        raise HTTPException(404, detail="客源线索不存在")
    return _lead(row, redaction_level)


# 客源婚况（varchar 档位）→ users.is_married（int）编码：0未知 1未婚 2离异 3丧偶。
# 只搬运能干净映射的取值（数字码或中文标签），无法识别时不写回（保持账号默认）。
_LEAD_MARRIAGE_TO_CODE = {"未婚": 1, "离异": 2, "丧偶": 3}


def _map_lead_marriage(value: str | None) -> int | None:
    if value is None:
        return None
    text_value = str(value).strip()
    if text_value.isdigit():
        code = int(text_value)
        return code if code in (1, 2, 3) else None
    return _LEAD_MARRIAGE_TO_CODE.get(text_value)


async def _table_columns(db: AsyncSession, table: str) -> frozenset[str]:
    rows = await db.execute(text("""SELECT COLUMN_NAME FROM information_schema.COLUMNS
        WHERE TABLE_SCHEMA = DATABASE() AND TABLE_NAME = :t"""), {"t": table})
    return frozenset(row[0] for row in rows.all())


async def _upsert_profile_row(
    db: AsyncSession, table: str, user_id: int, data: dict[str, Any], existing_cols: frozenset[str]
) -> None:
    """向 user_profile / user_auth 写一行资料（按 user_id 唯一键 upsert）。

    只写非 None 且列真实存在的字段（缺列静默跳过，保证「先发版后迁移」不报错）。
    表名为服务内部常量，不拼接用户输入。
    """
    cols = [c for c, v in data.items() if v is not None and c in existing_cols]
    if not cols:
        return
    params: dict[str, Any] = {c: data[c] for c in cols}
    params["user_id"] = user_id
    col_sql = ", ".join(cols)
    val_sql = ", ".join(f":{c}" for c in cols)
    upd_sql = ", ".join(f"{c} = VALUES({c})" for c in cols)
    await db.execute(text(
        f"""INSERT INTO {table} (user_id, {col_sql}) VALUES (:user_id, {val_sql})
            ON DUPLICATE KEY UPDATE {upd_sql}, updated_at = UTC_TIMESTAMP()"""), params)


async def convert_lead(
    db: AsyncSession,
    account_id: int,
    organization_id: int | None,
    matchmaker_user_id: int | None,
    lead_id: int,
    request: CustomerLeadConvertRequest,
    *,
    scope: str | None = None,
    scope_params: dict[str, Any] | None = None,
) -> tuple[CustomerLead, int, bool]:
    """Convert one admin lead into a pending member review atomically."""
    guard_where = ["id = :id"]
    guard_params: dict[str, Any] = {"id": lead_id}
    _merge_scope(guard_where, guard_params, scope, scope_params)
    if "deleted_at" in await _lead_columns(db):
        guard_where.append("deleted_at IS NULL")
    if not await db.scalar(text(f"SELECT id FROM customer_lead WHERE {' AND '.join(guard_where)} FOR UPDATE"), guard_params):
        raise HTTPException(404, detail="客源线索不存在")
    lead = await get_lead(db, lead_id, "full", scope=scope, scope_params=scope_params)
    if lead.converted_user_id:
        return lead, lead.converted_user_id, True
    if lead.status in ("LOST", "CLOSED", "CONVERTED"):
        raise HTTPException(409, detail="当前状态的客源不能入库")
    phone = normalize_contact(await _raw_contact(db, lead_id, "phone"))
    if not phone:
        raise HTTPException(409, detail="线索未登记手机号，无法生成账号，请先补充电话")
    existing = await db.execute(text("SELECT id FROM users WHERE phone = :phone LIMIT 1"), {"phone": phone})
    if existing.scalar():
        raise HTTPException(409, detail="该手机号已注册会员，不能重复入库")
    target_organization_id = lead.organization_id or organization_id
    target_matchmaker_id = lead.matchmaker_id or matchmaker_user_id
    try:
        result = await db.execute(text("""INSERT INTO users
            (openid, phone, phone_verified_at, nickname, gender, status, data_complete_rate, register_ip)
            VALUES (:openid, :phone, UTC_TIMESTAMP(), :nickname, :gender, 1, 0, '127.0.0.1')"""), {
            "openid": f"lead-convert-{lead_id}", "phone": phone, "nickname": lead.name, "gender": request.gender,
        })
        user_id = int(result.lastrowid)
        await ensure_default_user_role(db, user_id)
        # 客源「本人资料」搬运到新账号（对标站建档即带资料）：user_profile / user_auth 走 upsert，
        # users.birthday / is_married 属主表基础列直接补写。income（varchar 档位）与账号 decimal
        # 口径不兼容、不搬运；缺列静默跳过。
        from app.services.profile import _json_value  # 延迟导入避免循环依赖

        profile_cols = await _table_columns(db, "user_profile")
        auth_cols = await _table_columns(db, "user_auth")
        await _upsert_profile_row(db, "user_profile", user_id, {
            "height": lead.height, "weight": lead.weight, "hometown": lead.hometown,
            "residence": lead.residence, "ethnicity": lead.ethnicity, "house": lead.house,
            "car": lead.car, "smoking": lead.smoking, "drinking": lead.drinking,
            "occupation": lead.occupation,
            "photos": _json_value(lead.photos) if lead.photos else None,
        }, profile_cols)
        await _upsert_profile_row(db, "user_auth", user_id, {
            "real_name": lead.real_name, "education": lead.education, "school": lead.school,
            "company": lead.company, "job": lead.occupation,
        }, auth_cols)
        user_patch: dict[str, Any] = {}
        if lead.birth_date is not None:
            user_patch["birthday"] = lead.birth_date
        mapped_marriage = _map_lead_marriage(lead.marriage)
        if mapped_marriage is not None:
            user_patch["is_married"] = mapped_marriage
        if user_patch:
            patch_assign = ", ".join(f"{k} = :{k}" for k in user_patch)
            await db.execute(text(f"UPDATE users SET {patch_assign} WHERE id = :uid"), {**user_patch, "uid": user_id})
        await db.execute(text("""INSERT INTO matchmaker_member_review (user_id, status, reviewed_by)
            VALUES (:user_id, 'PENDING', NULL)
            ON DUPLICATE KEY UPDATE updated_at = UTC_TIMESTAMP()"""), {"user_id": user_id})
        if target_organization_id is not None or target_matchmaker_id is not None:
            await db.execute(text("""INSERT INTO resource_assignment
                (user_id, organization_id, matchmaker_id, source, status, assigned_by)
                VALUES (:user_id, :organization_id, :matchmaker_id, 'lead_convert', 1, :assigned_by)"""), {
                "user_id": user_id, "organization_id": target_organization_id,
                "matchmaker_id": target_matchmaker_id, "assigned_by": account_id,
            })
        if request.remark:
            await db.execute(text("""INSERT INTO matchmaker_admin_member_note (user_id, note, updated_by)
                VALUES (:user_id, :note, :updated_by)
                ON DUPLICATE KEY UPDATE note = VALUES(note), updated_by = VALUES(updated_by), updated_at = UTC_TIMESTAMP()"""),
                {"user_id": user_id, "note": request.remark, "updated_by": account_id})
        await db.execute(text("""UPDATE customer_lead
            SET status = 'CONVERTED', converted_user_id = :user_id, updated_at = UTC_TIMESTAMP()
            WHERE id = :lead_id AND converted_user_id IS NULL"""), {"user_id": user_id, "lead_id": lead_id})
        # 未入库期间在 customer_lead_preference 维护的择偶要求，入库时搬运至新账号的
        # user_partner_preference（INSERT..SELECT：线索未填过偏好则无行、自然跳过）。
        pref_cols = ", ".join(_LEAD_PREF_COLUMNS)
        await db.execute(text(
            f"INSERT INTO user_partner_preference (user_id, {pref_cols}) "
            f"SELECT :user_id, {pref_cols} FROM customer_lead_preference WHERE lead_id = :lead_id"
        ), {"user_id": user_id, "lead_id": lead_id})
        await db.execute(text("""INSERT INTO business_audit_log
            (actor_user_id, action, resource_type, resource_id, reason)
            VALUES (:actor, 'customer_lead.convert', 'customer_lead', :id, :reason)"""),
            {"actor": account_id, "id": lead_id, "reason": f"一键入库生成用户 {user_id}"})
        await db.commit()
    except IntegrityError as exc:
        await db.rollback()
        raise HTTPException(409, detail="客源入库失败：手机号或账号已存在") from exc
    return await get_lead(db, lead_id, REDACTION_LEVEL_STANDARD), user_id, False


async def _raw_contact(db: AsyncSession, lead_id: int, column: str) -> str | None:
    """重复性校验必须比对库里的真实联系方式，不能用脱敏后的展示值。"""
    if column not in ("phone", "wechat"):
        raise ValueError(f"不支持的联系方式字段: {column}")
    row = (await db.execute(text(f"SELECT {column} FROM customer_lead WHERE id = :id"), {"id": lead_id})).first()
    return row[0] if row else None


# 客源线索列表排序白名单；值全部服务端写死，不拼接用户输入。
_LEAD_ORDER_SQL = {
    "latest": "id DESC",
    "earliest": "id ASC",
}

# 「隐藏今日已跟进 / 已通话」落到当日存在记录的排除条件。
_LEAD_HIDE_TODAY_SQL = {
    "followed": """NOT EXISTS (SELECT 1 FROM customer_lead_follow_up fu
        WHERE fu.lead_id = customer_lead.id AND fu.created_at >= CURDATE()
          AND fu.created_at < DATE_ADD(CURDATE(), INTERVAL 1 DAY))""",
    "called": """NOT EXISTS (SELECT 1 FROM customer_lead_call_note cn
        WHERE cn.lead_id = customer_lead.id AND cn.created_at >= CURDATE()
          AND cn.created_at < DATE_ADD(CURDATE(), INTERVAL 1 DAY))""",
}

async def _lead_columns(db: AsyncSession) -> frozenset[str]:
    rows = await db.execute(text("""SELECT COLUMN_NAME FROM information_schema.COLUMNS
        WHERE TABLE_SCHEMA = DATABASE() AND TABLE_NAME = 'customer_lead'"""))
    return frozenset(row[0] for row in rows.all())


def _split_multi(raw: str | None, limit: int = 50) -> list[str]:
    """多选参数按英文逗号拆成去重取值列表。"""
    if not raw:
        return []
    return [item.strip() for item in dict.fromkeys(raw.split(",")) if item.strip()][:limit]


def _placeholders(prefix: str, values: list[str], params: dict[str, Any]) -> str:
    """为多值 IN 生成命名占位符并写回 params，避免拼接用户输入。"""
    names = []
    for index, value in enumerate(values):
        key = f"{prefix}_{index}"
        params[key] = value
        names.append(f":{key}")
    return ", ".join(names)


async def list_leads(
    db: AsyncSession,
    page: int,
    page_size: int,
    status: str | None,
    source: str | None,
    matchmaker_id: int | None,
    search: str | None,
    redaction_level: str = REDACTION_LEVEL_STANDARD,
    *,
    intention_level: int | None = None,
    gender: int | None = None,
    start_date: str | None = None,
    end_date: str | None = None,
    assign_start: str | None = None,
    assign_end: str | None = None,
    promoter_id: int | None = None,
    xiangqin_status: int | None = None,
    bind_mode: str | None = None,
    wechat_added: int | None = None,
    age_min: int | None = None,
    age_max: int | None = None,
    height_min: int | None = None,
    height_max: int | None = None,
    marriage: str | None = None,
    education: str | None = None,
    occupation: str | None = None,
    residence: str | None = None,
    hometown: str | None = None,
    hide_followed_today: bool | None = None,
    hide_called_today: bool | None = None,
    has_phone: bool | None = None,
    has_wechat: bool | None = None,
    assigned: bool | None = None,
    audit_status: str | None = None,
    stored: bool | None = None,
    order: str | None = None,
    scope: str | None = None,
    scope_params: dict[str, Any] | None = None,
) -> CustomerLeadPage:
    """客源线索分页列表。

    新增筛选项依赖 `migrations/filters/20261011_01_admin_filter_columns_up.sql`
    补齐的 customer_lead 列；列不存在时对应条件静默跳过，保证「先发版后迁移」
    不会让列表整页 500。

    `scope` / `scope_params` 由路由层用 `CurrentMatchmakerAdmin.scope_clause`（D-16）
    生成并下传；未传表示不做约束（超管或服务层直调）。
    """
    columns = await _lead_columns(db)
    where = ["1=1"]
    params: dict[str, Any] = {"limit": page_size, "offset": (page - 1) * page_size}
    _merge_scope(where, params, scope, scope_params)
    if "deleted_at" in columns:
        where.append("deleted_at IS NULL")
    if status:
        where.append("status = :status")
        params["status"] = status
    if source:
        where.append("source = :source")
        params["source"] = source
    if matchmaker_id is not None:
        where.append("matchmaker_id = :matchmaker_id")
        params["matchmaker_id"] = matchmaker_id
    if search:
        where.append("(name LIKE CONCAT('%', :search, '%') OR phone LIKE CONCAT('%', :search, '%') OR wechat LIKE CONCAT('%', :search, '%'))")
        params["search"] = search

    if intention_level is not None:
        where.append("intention_level = :intention_level")
        params["intention_level"] = intention_level
    if gender is not None and "gender" in columns:
        where.append("gender = :gender")
        params["gender"] = gender
    if promoter_id is not None:
        where.append("promoter_id = :promoter_id")
        params["promoter_id"] = promoter_id
    # B11③ 客源线索缺失筛选：待分派 / 审核状态 / 入库情况。
    # matchmaker_id、audit_status、converted_user_id 均为 customer_lead 基础列（LEAD_SELECT 已引用）。
    if assigned is True:
        where.append("matchmaker_id IS NOT NULL")
    elif assigned is False:
        where.append("matchmaker_id IS NULL")
    if audit_status:
        where.append("audit_status = :audit_status")
        params["audit_status"] = audit_status
    if stored is True:
        where.append("converted_user_id IS NOT NULL")
    elif stored is False:
        where.append("converted_user_id IS NULL")
    if start_date:
        where.append("created_at >= :start_date")
        params["start_date"] = start_date
    if end_date:
        # 结束日按「含当日」处理。
        where.append("created_at < DATE_ADD(:end_date, INTERVAL 1 DAY)")
        params["end_date"] = end_date
    if hide_followed_today:
        where.append(_LEAD_HIDE_TODAY_SQL["followed"])
    if hide_called_today:
        where.append(_LEAD_HIDE_TODAY_SQL["called"])
    if has_phone is True:
        where.append("phone IS NOT NULL AND TRIM(phone) <> ''")
    elif has_phone is False:
        where.append("(phone IS NULL OR TRIM(phone) = '')")
    if has_wechat is True:
        where.append("wechat IS NOT NULL AND TRIM(wechat) <> ''")
    elif has_wechat is False:
        where.append("(wechat IS NULL OR TRIM(wechat) = '')")

    # 以下筛选项依赖新增列，缺列时跳过。
    def has(name: str) -> bool:
        return name in columns

    if has("assigned_at"):
        if assign_start:
            where.append("assigned_at >= :assign_start")
            params["assign_start"] = assign_start
        if assign_end:
            where.append("assigned_at < DATE_ADD(:assign_end, INTERVAL 1 DAY)")
            params["assign_end"] = assign_end
    if xiangqin_status is not None and has("xiangqin_status"):
        where.append("xiangqin_status = :xiangqin_status")
        params["xiangqin_status"] = xiangqin_status
    if bind_mode and has("bind_mode"):
        where.append("bind_mode = :bind_mode")
        params["bind_mode"] = bind_mode
    if wechat_added is not None and has("wechat_added"):
        where.append("wechat_added = :wechat_added")
        params["wechat_added"] = wechat_added
    if age_min is not None and has("age"):
        where.append("age >= :age_min")
        params["age_min"] = age_min
    if age_max is not None and has("age"):
        where.append("age <= :age_max")
        params["age_max"] = age_max
    if height_min is not None and has("height"):
        where.append("height >= :height_min")
        params["height_min"] = height_min
    if height_max is not None and has("height"):
        where.append("height <= :height_max")
        params["height_max"] = height_max
    for column, raw in (("marriage", marriage), ("education", education), ("occupation", occupation)):
        values = _split_multi(raw)
        if values and has(column):
            where.append(f"{column} IN ({_placeholders(column, values, params)})")
    for column, raw in (("residence", residence), ("hometown", hometown)):
        if raw and has(column):
            where.append(f"{column} LIKE CONCAT('%', :{column}, '%')")
            params[column] = raw

    clause = " AND ".join(where)
    order_sql = _LEAD_ORDER_SQL.get(order or "latest", _LEAD_ORDER_SQL["latest"])
    select_sql = _lead_select_for(columns)
    rows = await db.execute(text(f"{select_sql} WHERE {clause} ORDER BY {order_sql} LIMIT :limit OFFSET :offset"), params)
    count = await db.execute(text(f"SELECT COUNT(*) FROM customer_lead WHERE {clause}"), {k: v for k, v in params.items() if k not in ("limit", "offset")})
    total = int(count.scalar() or 0)
    return CustomerLeadPage(items=[_lead(row, redaction_level) for row in rows.mappings().all()], page=page, page_size=page_size, total=total, has_more=page * page_size < total)


async def update_lead(
    db: AsyncSession,
    account_id: int,
    lead_id: int,
    request: CustomerLeadUpdate,
    redaction_level: str = REDACTION_LEVEL_STANDARD,
    *,
    scope: str | None = None,
    scope_params: dict[str, Any] | None = None,
) -> CustomerLead:
    current = await get_lead(db, lead_id, redaction_level, scope=scope, scope_params=scope_params)
    from app.services.profile import _calculate_age, _json_value  # 延迟导入避免循环依赖
    columns = await _lead_columns(db)
    values = request.model_dump(exclude_unset=True)
    # photos 存 JSON；birth_date 有值时同步派生 raw age 列。
    if values.get("photos") is not None:
        values["photos"] = _json_value(values["photos"])
    if values.get("birth_date") is not None and "age" in columns:
        values["age"] = _calculate_age(values["birth_date"])
    # 缺列降级：库中不存在的 gender / 资料列 / age 从写入字段剔除，避免 Unknown column 500（先发版后迁移）。
    for col in ("gender", *_LEAD_PROFILE_COLUMNS, "age"):
        if col in values and col not in columns:
            values.pop(col)
    # 掩码值不写回：表单预填的是脱敏后的展示值，原样落库会把真实微信号毁成 ***
    if is_masked_value(values.get("wechat")):
        values.pop("wechat")
    target_status = values.get("status") or current.status
    # 目标状态仍为有效线索时，联系方式不得与其它有效线索重复；改为 LOST/CLOSED 后生成列会置空。
    if target_status not in ("LOST", "CLOSED"):
        target_phone = normalize_contact(values["phone"]) if "phone" in values else normalize_contact(await _raw_contact(db, lead_id, "phone"))
        target_wechat = normalize_contact(values["wechat"]) if "wechat" in values else normalize_contact(await _raw_contact(db, lead_id, "wechat"))
        await ensure_contact_available(db, target_phone, target_wechat, exclude_lead_id=lead_id)
    assignments = ", ".join(f"{key} = :{key}" for key in values)
    # 掩码回显字段被全部丢弃时 assignments 为空，直接拼会产出 ``SET , updated_at`` 病态 SQL；此时只刷时间戳。
    set_clause = f"{assignments}, " if assignments else ""
    # 写路径同样带上作用域谓词（D-16）：读可见校验与写之间若被改派，UPDATE 也不越权落库。
    update_where = ["id = :id"]
    update_params: dict[str, Any] = {**values, "id": lead_id}
    _merge_scope(update_where, update_params, scope, scope_params)
    try:
        await db.execute(text(f"UPDATE customer_lead SET {set_clause}updated_at = UTC_TIMESTAMP() WHERE {' AND '.join(update_where)}"), update_params)
        await db.execute(text("INSERT INTO business_audit_log (actor_user_id, action, resource_type, resource_id) VALUES (:actor, 'customer_lead.update', 'customer_lead', :id)"), {"actor": account_id, "id": lead_id})
        await db.commit()
    except IntegrityError:
        await db.rollback()
        raise_duplicate_contact()
    return await get_lead(db, lead_id, redaction_level, scope=scope, scope_params=scope_params)


# 未入库客源择偶要求可写列（与 C 端 PreferenceUpdateRequest 口径一致，主键为 lead_id）。
_LEAD_PREF_COLUMNS = [
    "age_min", "age_max", "height_min", "height_max", "education_min", "income_min",
    "marriage_status", "preferred_province_code", "preferred_city_codes",
    "accept_long_distance", "accept_cross_province", "housing_requirement",
    "smoking_requirement", "drinking_requirement", "dating_goal", "meeting_pace",
    "children_intention", "extra_requirement",
]


async def _get_lead_preference_dict(db: AsyncSession, lead_id: int) -> dict[str, Any]:
    """读取未入库客源的择偶要求（customer_lead_preference），无行时返回空默认值。

    输出形状对齐 PreferenceResponse：user_id 固定为 0（未入库无账号），
    preferred_city_codes/布尔列按 C 端口径反解。多余键（id/lead_id/时间戳）由
    PreferenceResponse 忽略。
    """
    from app.services.profile import _json_list  # 延迟导入避免模块级循环依赖

    result = await db.execute(
        text("SELECT * FROM customer_lead_preference WHERE lead_id = :lead_id"), {"lead_id": lead_id}
    )
    row = result.mappings().first()
    if not row:
        return {
            "user_id": 0, "preferred_city_codes": [],
            "accept_long_distance": False, "accept_cross_province": False,
            **{c: None for c in _LEAD_PREF_COLUMNS
               if c not in ("preferred_city_codes", "accept_long_distance", "accept_cross_province")},
        }
    data = dict(row)
    data["user_id"] = 0
    data["preferred_city_codes"] = _json_list(data["preferred_city_codes"])
    data["accept_long_distance"] = bool(data["accept_long_distance"])
    data["accept_cross_province"] = bool(data["accept_cross_province"])
    return data


async def _upsert_lead_preference(db: AsyncSession, lead_id: int, request: PreferenceUpdateRequest) -> None:
    """部分更新未入库客源择偶要求（与 profile.update_preferences 同口径，但不涉及完整度/画像）。"""
    from app.services.profile import _json_value  # 延迟导入避免模块级循环依赖

    values = request.model_dump(exclude_unset=True)
    if "preferred_city_codes" in values:
        values["preferred_city_codes"] = _json_value(values["preferred_city_codes"])
    values = {k: (int(v) if isinstance(v, bool) else v) for k, v in values.items()}
    exists = (await db.execute(
        text("SELECT id FROM customer_lead_preference WHERE lead_id = :lead_id"), {"lead_id": lead_id}
    )).scalar()
    if exists and values:
        assignments = ", ".join(f"{k} = :{k}" for k in values)
        await db.execute(
            text(f"UPDATE customer_lead_preference SET {assignments}, updated_at = UTC_TIMESTAMP() WHERE lead_id = :lead_id"),
            {**values, "lead_id": lead_id},
        )
    elif not exists:
        columns = ["lead_id", *values]
        placeholders = ", ".join(f":{c}" for c in columns)
        await db.execute(
            text(f"INSERT INTO customer_lead_preference ({', '.join(columns)}) VALUES ({placeholders})"),
            {"lead_id": lead_id, **values},
        )


async def get_lead_preferences(
    db: AsyncSession,
    lead_id: int,
    *,
    scope: str | None = None,
    scope_params: dict[str, Any] | None = None,
) -> LeadPreferenceResponse:
    """读取客源线索的择偶要求。

    先走 get_lead 做作用域与存在校验（D-16）。已入库（有 converted_user_id）读账号那份
    user_partner_preference；未入库读 customer_lead_preference 暂存表（无行则返回空默认值）。
    两种情况 editable 均为 True（未入库现已可维护）。
    """
    lead = await get_lead(db, lead_id, scope=scope, scope_params=scope_params)
    if lead.converted_user_id:
        # 延迟导入避免与 profile 服务形成模块级循环依赖。
        from app.services.profile import get_preferences

        data = await get_preferences(db, lead.converted_user_id)
        return LeadPreferenceResponse(
            lead_id=lead_id, converted_user_id=lead.converted_user_id, editable=True, preference=data
        )
    data = await _get_lead_preference_dict(db, lead_id)
    return LeadPreferenceResponse(
        lead_id=lead_id, converted_user_id=None, editable=True, preference=data
    )


async def update_lead_preferences(
    db: AsyncSession,
    account_id: int,
    lead_id: int,
    request: PreferenceUpdateRequest,
    *,
    scope: str | None = None,
    scope_params: dict[str, Any] | None = None,
) -> LeadPreferenceResponse:
    """写入客源线索的择偶要求。

    已入库则复用 C 端 update_preferences（含完整度重算与画像 revision 入队），与会员
    自助编辑同口径；未入库写 customer_lead_preference 暂存表，入库时再搬运至账号。
    两种情况均追加审计 action='customer_lead.preference_update'。
    """
    lead = await get_lead(db, lead_id, scope=scope, scope_params=scope_params)
    if lead.converted_user_id:
        from app.services.profile import update_preferences

        data = await update_preferences(db, lead.converted_user_id, request)
    else:
        await _upsert_lead_preference(db, lead_id, request)
        data = await _get_lead_preference_dict(db, lead_id)
    await db.execute(
        text("INSERT INTO business_audit_log (actor_user_id, action, resource_type, resource_id) "
             "VALUES (:actor, 'customer_lead.preference_update', 'customer_lead', :id)"),
        {"actor": account_id, "id": lead_id},
    )
    await db.commit()
    return LeadPreferenceResponse(
        lead_id=lead_id, converted_user_id=lead.converted_user_id, editable=True, preference=data
    )


async def assign_lead(
    db: AsyncSession,
    account_id: int,
    lead_id: int,
    request: CustomerLeadAssignment,
    redaction_level: str = REDACTION_LEVEL_STANDARD,
    *,
    scope: str | None = None,
    scope_params: dict[str, Any] | None = None,
) -> CustomerLead:
    await get_lead(db, lead_id, scope=scope, scope_params=scope_params)
    await _validate_owner(db, request)
    # 只写调用方显式给出的归属人：只改红娘时不能把门店冲成 NULL。
    values = request.model_dump(exclude_unset=True)
    assign_where = ["id = :id"]
    assign_params: dict[str, Any] = {**values, "id": lead_id}
    _merge_scope(assign_where, assign_params, scope, scope_params)
    set_parts = [f"{key} = :{key}" for key in values]
    # 归属变更同步刷新分派时间；显式清空归属时一并清空，避免「无归属却有分派时间」。
    if "matchmaker_id" in values and "assigned_at" in await _lead_columns(db):
        set_parts.append("assigned_at = CASE WHEN :matchmaker_id IS NULL THEN NULL ELSE UTC_TIMESTAMP() END")
    set_clause = ", ".join(set_parts)
    await db.execute(text(f"UPDATE customer_lead SET {set_clause}, updated_at = UTC_TIMESTAMP() WHERE {' AND '.join(assign_where)}"), assign_params)
    await db.execute(text("INSERT INTO business_audit_log (actor_user_id, action, resource_type, resource_id) VALUES (:actor, 'customer_lead.assign', 'customer_lead', :id)"), {"actor": account_id, "id": lead_id})
    await db.commit()
    return await get_lead(db, lead_id, redaction_level, scope=scope, scope_params=scope_params)


async def auto_assign_lead(
    db: AsyncSession,
    account_id: int,
    lead_id: int,
    *,
    force: bool = False,
    redaction_level: str = REDACTION_LEVEL_STANDARD,
    scope: str | None = None,
    scope_params: dict[str, Any] | None = None,
) -> CustomerLeadCreateResult:
    """对已存在的客源按分派策略重算归属红娘。

    默认仅对当前未分派的客源生效；force=true 即使已有归属也重算并覆盖。
    选不到人时不改动现有归属，返回 assigned=false，让前端可提示。
    """
    lead = await get_lead(db, lead_id, redaction_level, scope=scope, scope_params=scope_params)
    if lead.matchmaker_id is not None and not force:
        return CustomerLeadCreateResult(**lead.model_dump(), assigned=True)
    matchmaker_id = await resolve_assigned_matchmaker(db, auto_assign=True)
    if not matchmaker_id:
        return CustomerLeadCreateResult(**lead.model_dump(), assigned=False)
    assign_where = ["id = :id"]
    assign_params: dict[str, Any] = {"id": lead_id, "matchmaker_id": matchmaker_id}
    _merge_scope(assign_where, assign_params, scope, scope_params)
    assigned_at_sql = ", assigned_at = UTC_TIMESTAMP()" if "assigned_at" in await _lead_columns(db) else ""
    await db.execute(text(f"UPDATE customer_lead SET matchmaker_id = :matchmaker_id{assigned_at_sql}, updated_at = UTC_TIMESTAMP() WHERE {' AND '.join(assign_where)}"), assign_params)
    await db.execute(text("INSERT INTO business_audit_log (actor_user_id, action, resource_type, resource_id) VALUES (:actor, 'customer_lead.auto_assign', 'customer_lead', :id)"), {"actor": account_id, "id": lead_id})
    await db.commit()
    updated = await get_lead(db, lead_id, redaction_level, scope=scope, scope_params=scope_params)
    return CustomerLeadCreateResult(**updated.model_dump(), assigned=True)


async def add_follow_up(
    db: AsyncSession,
    account_id: int,
    lead_id: int,
    request: CustomerLeadFollowUpCreate,
    *,
    scope: str | None = None,
    scope_params: dict[str, Any] | None = None,
) -> CustomerLeadFollowUp:
    await get_lead(db, lead_id, scope=scope, scope_params=scope_params)
    result = await db.execute(text("""INSERT INTO customer_lead_follow_up
        (lead_id, method, content, intention_level, next_follow_at, created_by)
        VALUES (:lead_id, :method, :content, :intention_level, :next_follow_at, :created_by)"""), {**request.model_dump(), "lead_id": lead_id, "created_by": account_id})
    follow_id = int(result.lastrowid)
    await db.execute(text("UPDATE customer_lead SET status = CASE WHEN status = 'NEW' THEN 'CONTACTED' ELSE status END, next_follow_at = :next_follow_at, updated_at = UTC_TIMESTAMP() WHERE id = :id"), {"next_follow_at": request.next_follow_at, "id": lead_id})
    await db.commit()
    row = (await db.execute(text("SELECT id, lead_id, method, content, intention_level, next_follow_at, created_by, created_at FROM customer_lead_follow_up WHERE id = :id"), {"id": follow_id})).mappings().one()
    return CustomerLeadFollowUp(**dict(row))


_LEAD_FU_BASE_COLS = "id, lead_id, method, content, intention_level, next_follow_at, created_by, created_at"
_LEAD_FU_MEDIA_COLS = ("images", "voice_url", "voice_duration_sec")


def _lead_fu_media_select(cols: frozenset[str]) -> str:
    """根据实际存在的附件列拼接 SELECT 片段（缺列降级，支持先发版后迁移）。"""
    return "".join(f", {col}" for col in _LEAD_FU_MEDIA_COLS if col in cols)


def _build_lead_follow_up(row: dict[str, Any]) -> CustomerLeadFollowUp:
    """行 -> 响应模型。images 反序列化为 list，voice_duration_sec 转 int，缺列时走默认。"""
    payload = dict(row)
    if "images" in payload:
        payload["images"] = _parse_follow_up_images(payload["images"])
    if payload.get("voice_duration_sec") is not None:
        payload["voice_duration_sec"] = int(payload["voice_duration_sec"])
    return CustomerLeadFollowUp(**payload)


def _sign_lead_follow_up(item: CustomerLeadFollowUp, viewer: int) -> CustomerLeadFollowUp:
    """响应层重签：DB 存原始 URL，返回给当前管理员的附件 URL 附短期签名（TTL 300s）。"""
    item.images = [
        sign_media_url(url, category=MEDIA_CATEGORY_FOLLOW_UP, viewer=viewer) for url in item.images
    ]
    if item.voice_url:
        item.voice_url = sign_media_url(item.voice_url, category=MEDIA_CATEGORY_FOLLOW_UP, viewer=viewer)
    return item


async def create_follow_up_with_media(
    db: AsyncSession,
    account_id: int,
    lead_id: int,
    method: str,
    content: str,
    intention_level: int | None,
    next_follow_at: datetime | None,
    images: list[UploadFile] | None,
    voice: UploadFile | None,
    voice_duration_sec: int | None,
    *,
    scope: str | None = None,
    scope_params: dict[str, Any] | None = None,
) -> CustomerLeadFollowUp:
    """新增客源跟进记录，支持文字 + 图片（转 webp，最多 9 张）+ 录音。

    与会员侧 create_follow_up_with_media 完全同构：图片存
    ``{upload_dir}/{lead_id}/follow-up-img-{uuid}.webp``，录音存
    ``follow-up-voice-{uuid}.{ext}``；URL 形如 ``/storage/uploads/{lead_id}/...``
    （命中 MEDIA_CATEGORY_FOLLOW_UP 的私有签名路径）。越权（线索不在数据范围内）
    由 get_lead 抛 404，不落库。附件列缺失时静默降级为仅存文字（先发版后迁移）。
    """
    await get_lead(db, lead_id, scope=scope, scope_params=scope_params)
    if method not in {"PHONE", "WECHAT", "VISIT", "OTHER"}:
        raise HTTPException(422, detail="跟进方式必须为 PHONE/WECHAT/VISIT/OTHER")
    normalized_content = (content or "").strip()
    image_files = [item for item in (images or []) if item is not None and (item.filename or "").strip()]
    has_voice = voice is not None and bool((voice.filename or "").strip()) and voice.size != 0
    if not normalized_content and not image_files and not has_voice:
        raise HTTPException(422, detail="跟进文字与图片/录音至少提供一项")
    if len(image_files) > FOLLOW_UP_MAX_IMAGES:
        raise HTTPException(422, detail=f"图片最多{FOLLOW_UP_MAX_IMAGES}张")

    directory = Path(settings.upload_dir) / str(lead_id)
    image_urls: list[str] = []
    for image in image_files:
        data = await _fu_read_limited(image, FOLLOW_UP_IMAGE_MAX_BYTES)
        webp_data, _thumb = await asyncio.to_thread(_fu_image_outputs, data)
        name = uuid.uuid4().hex
        target = directory / f"follow-up-img-{name}.webp"
        await _fu_write_bytes(target, webp_data)
        image_urls.append(f"/storage/uploads/{lead_id}/{target.name}")

    voice_url: str | None = None
    if has_voice:
        voice_data = await _fu_read_limited(voice, FOLLOW_UP_VOICE_MAX_BYTES)
        suffix = _voice_suffix(voice)
        # 时长以服务端解析为准，前端传值仅在容器无法识别时兜底。
        voice_duration_sec = parse_audio_duration(voice_data) or voice_duration_sec
        name = uuid.uuid4().hex
        target = directory / f"follow-up-voice-{name}.{suffix}"
        await _fu_write_bytes(target, voice_data)
        voice_url = f"/storage/uploads/{lead_id}/{target.name}"

    cols = await _table_columns(db, "customer_lead_follow_up")
    insert_cols = ["lead_id", "method", "content", "intention_level", "next_follow_at", "created_by"]
    params: dict[str, Any] = {
        "lead_id": lead_id, "method": method, "content": normalized_content,
        "intention_level": intention_level, "next_follow_at": next_follow_at, "created_by": account_id,
    }
    if set(_LEAD_FU_MEDIA_COLS) <= cols:
        insert_cols += list(_LEAD_FU_MEDIA_COLS)
        params["images"] = json.dumps(image_urls, ensure_ascii=False) if image_urls else None
        params["voice_url"] = voice_url
        params["voice_duration_sec"] = voice_duration_sec if (voice_url and voice_duration_sec) else None
    col_sql = ", ".join(f"`{c}`" for c in insert_cols)
    val_sql = ", ".join(f":{c}" for c in insert_cols)
    result = await db.execute(text(f"INSERT INTO customer_lead_follow_up ({col_sql}) VALUES ({val_sql})"), params)
    follow_id = int(result.lastrowid)
    await db.execute(text("UPDATE customer_lead SET status = CASE WHEN status = 'NEW' THEN 'CONTACTED' ELSE status END, next_follow_at = :next_follow_at, updated_at = UTC_TIMESTAMP() WHERE id = :id"), {"next_follow_at": next_follow_at, "id": lead_id})
    await db.execute(text("""INSERT INTO business_audit_log (actor_user_id, action, resource_type, resource_id, after_json)
        VALUES (:actor, 'customer_lead.follow_up.create_media', 'customer_lead_follow_up', :rid, :after)"""),
        {"actor": account_id, "rid": follow_id, "after": json.dumps({"images": image_urls, "voice_url": voice_url}, ensure_ascii=False)})
    await db.commit()
    row = (await db.execute(text(f"SELECT {_LEAD_FU_BASE_COLS}{_lead_fu_media_select(cols)} FROM customer_lead_follow_up WHERE id = :id"), {"id": follow_id})).mappings().one()
    return _sign_lead_follow_up(_build_lead_follow_up(dict(row)), account_id)


async def list_follow_ups(
    db: AsyncSession,
    lead_id: int,
    page: int,
    page_size: int,
    *,
    viewer: int | None = None,
    scope: str | None = None,
    scope_params: dict[str, Any] | None = None,
) -> list[CustomerLeadFollowUp]:
    await get_lead(db, lead_id, scope=scope, scope_params=scope_params)
    cols = await _table_columns(db, "customer_lead_follow_up")
    rows = await db.execute(text(f"""SELECT {_LEAD_FU_BASE_COLS}{_lead_fu_media_select(cols)}
        FROM customer_lead_follow_up WHERE lead_id = :lead_id ORDER BY id DESC LIMIT :limit OFFSET :offset"""), {"lead_id": lead_id, "limit": page_size, "offset": (page - 1) * page_size})
    items = [_build_lead_follow_up(dict(row)) for row in rows.mappings().all()]
    if viewer is not None:
        items = [_sign_lead_follow_up(item, viewer) for item in items]
    return items


async def lead_statistics(
    db: AsyncSession,
    *,
    scope: str | None = None,
    scope_params: dict[str, Any] | None = None,
) -> CustomerLeadStatistics:
    where = ["1=1"]
    params: dict[str, Any] = {}
    _merge_scope(where, params, scope, scope_params)
    if "deleted_at" in await _lead_columns(db):
        where.append("deleted_at IS NULL")
    row = (await db.execute(text(f"""SELECT COUNT(*) total, SUM(status = 'NEW') new_count, SUM(status = 'CONTACTED') contacted_count,
        SUM(status = 'INTENDED') intended_count, SUM(status = 'CONVERTED') converted_count, SUM(status = 'LOST') lost_count
        FROM customer_lead WHERE {' AND '.join(where)}"""), params)).mappings().one()
    return CustomerLeadStatistics(**{key: int(row[key] or 0) for key in ("total", "new_count", "contacted_count", "intended_count", "converted_count", "lost_count")})


# ───────────── A13 客源跟进全览 / 数据报表 ─────────────

async def list_follow_up_records(
    db: AsyncSession,
    page: int,
    page_size: int,
    *,
    matchmaker_id: int | None = None,
    promoter_id: int | None = None,
    intention_level: int | None = None,
    start_date: str | None = None,
    end_date: str | None = None,
    search: str | None = None,
    scope: str | None = None,
    scope_params: dict[str, Any] | None = None,
) -> CustomerLeadFollowUpRecordPage:
    """全局客源跟进记录列表（A13）：代替前端“拉全部线索→逐条跟进”的 N+1。

    归属 / 软删除均在 `customer_lead`，用 EXISTS 子查询收敛，既复用 D-16
    作用域谓词（列名不带别名，作用于子查询内的 customer_lead），又与列表明细同口径。
    """
    where = ["1=1"]
    params: dict[str, Any] = {"limit": page_size, "offset": (page - 1) * page_size}
    lead_where: list[str] = []
    lead_params: dict[str, Any] = {}
    if "deleted_at" in await _lead_columns(db):
        lead_where.append("customer_lead.deleted_at IS NULL")
    if scope:
        lead_where.append(f"({scope})")
        lead_params.update(scope_params or {})
    if matchmaker_id is not None:
        lead_where.append("customer_lead.matchmaker_id = :matchmaker_id")
        params["matchmaker_id"] = matchmaker_id
    if promoter_id is not None:
        lead_where.append("customer_lead.promoter_id = :promoter_id")
        params["promoter_id"] = promoter_id
    if intention_level is not None:
        where.append("fu.intention_level = :intention_level")
        params["intention_level"] = intention_level
    if start_date:
        where.append("fu.created_at >= :start")
        params["start"] = start_date
    if end_date:
        where.append("fu.created_at < DATE_ADD(:end, INTERVAL 1 DAY)")
        params["end"] = end_date
    if search:
        where.append("(fu.content LIKE CONCAT('%', :search, '%') OR cl.name LIKE CONCAT('%', :search, '%'))")
        params["search"] = search
    if lead_where:
        where.append("EXISTS (SELECT 1 FROM customer_lead WHERE customer_lead.id = fu.lead_id AND " + " AND ".join(lead_where) + ")")
    params.update(lead_params)
    clause = " AND ".join(where)
    count = int((await db.execute(text(f"SELECT COUNT(*) FROM customer_lead_follow_up fu JOIN customer_lead cl ON cl.id = fu.lead_id WHERE {clause}"), params)).scalar() or 0)
    rows = (await db.execute(text(f"""SELECT fu.id, fu.lead_id, cl.name AS lead_name, cl.source AS lead_source, cl.status AS lead_status,
        fu.method, fu.content, fu.intention_level, fu.next_follow_at, cl.matchmaker_id, cl.promoter_id, fu.created_by, fu.created_at
        FROM customer_lead_follow_up fu JOIN customer_lead cl ON cl.id = fu.lead_id
        WHERE {clause} ORDER BY fu.id DESC LIMIT :limit OFFSET :offset"""), params)).mappings().all()
    return CustomerLeadFollowUpRecordPage(
        items=[CustomerLeadFollowUpRecord(**dict(row)) for row in rows],
        page=page, page_size=page_size, total=count, has_more=page * page_size < count,
    )


async def lead_follow_statistics(
    db: AsyncSession,
    *,
    start_date: str | None = None,
    end_date: str | None = None,
    matchmaker_id: int | None = None,
    scope: str | None = None,
    scope_params: dict[str, Any] | None = None,
) -> CustomerLeadFollowBucket:
    """跟进全览顶部时间卡片计数（A13）：今天/昨天/最近3天/本周/上周/本月/上月。"""
    where = ["1=1"]
    params: dict[str, Any] = {}
    lead_where: list[str] = []
    if "deleted_at" in await _lead_columns(db):
        lead_where.append("customer_lead.deleted_at IS NULL")
    if scope:
        lead_where.append(f"({scope})")
        params.update(scope_params or {})
    if matchmaker_id is not None:
        lead_where.append("customer_lead.matchmaker_id = :matchmaker_id")
        params["matchmaker_id"] = matchmaker_id
    if start_date:
        where.append("fu.created_at >= :start")
        params["start"] = start_date
    if end_date:
        where.append("fu.created_at < DATE_ADD(:end, INTERVAL 1 DAY)")
        params["end"] = end_date
    if lead_where:
        where.append("EXISTS (SELECT 1 FROM customer_lead WHERE customer_lead.id = fu.lead_id AND " + " AND ".join(lead_where) + ")")
    clause = " AND ".join(where)
    row = (await db.execute(text(f"""SELECT
        SUM(fu.created_at >= CURDATE() AND fu.created_at < DATE_ADD(CURDATE(), INTERVAL 1 DAY)) today,
        SUM(fu.created_at >= DATE_SUB(CURDATE(), INTERVAL 1 DAY) AND fu.created_at < CURDATE()) yesterday,
        SUM(fu.created_at >= DATE_SUB(CURDATE(), INTERVAL 3 DAY)) last_3_days,
        SUM(fu.created_at >= DATE_SUB(CURDATE(), INTERVAL WEEKDAY(CURDATE()) DAY)) this_week,
        SUM(fu.created_at >= DATE_SUB(CURDATE(), INTERVAL (WEEKDAY(CURDATE()) + 7) DAY) AND fu.created_at < DATE_SUB(CURDATE(), INTERVAL WEEKDAY(CURDATE()) DAY)) last_week,
        SUM(fu.created_at >= DATE_FORMAT(CURDATE(), '%Y-%m-01')) this_month,
        SUM(fu.created_at >= DATE_SUB(DATE_FORMAT(CURDATE(), '%Y-%m-01'), INTERVAL 1 MONTH) AND fu.created_at < DATE_FORMAT(CURDATE(), '%Y-%m-01')) last_month
        FROM customer_lead_follow_up fu WHERE {clause}"""), params)).mappings().one()
    return CustomerLeadFollowBucket(**{key: int(row[key] or 0) for key in ("today", "yesterday", "last_3_days", "this_week", "last_week", "this_month", "last_month")})


async def matchmaker_follow_stats(
    db: AsyncSession,
    *,
    matchmaker_id: int | None = None,
    scope: str | None = None,
    scope_params: dict[str, Any] = None,
) -> list[MatchmakerFollowStat]:
    """按红娘的客源跟进统计（A13）：名下客源总数 / 跟进总条数 / 本月 / 从未跟进。"""
    if scope_params is None:
        scope_params = {}
    where = ["cl.matchmaker_id IS NOT NULL"]
    params: dict[str, Any] = {}
    if "deleted_at" in await _lead_columns(db):
        where.append("cl.deleted_at IS NULL")
    if scope:
        where.append(f"({scope})")
        params.update(scope_params)
    if matchmaker_id is not None:
        where.append("cl.matchmaker_id = :matchmaker_id")
        params["matchmaker_id"] = matchmaker_id
    clause = " AND ".join(where)
    rows = (await db.execute(text(f"""SELECT cl.matchmaker_id AS matchmaker_id,
        COALESCE(NULLIF(u.nickname, ''), u.phone, CONCAT('红娘', cl.matchmaker_id)) AS matchmaker_name,
        COUNT(DISTINCT cl.id) AS lead_count,
        COUNT(fu.id) AS total_count,
        SUM(fu.created_at >= DATE_FORMAT(CURDATE(), '%Y-%m-01')) AS month_count,
        COUNT(DISTINCT CASE WHEN NOT EXISTS (SELECT 1 FROM customer_lead_follow_up fx WHERE fx.lead_id = cl.id) THEN cl.id END) AS unfollowed_count
        FROM customer_lead cl
        LEFT JOIN customer_lead_follow_up fu ON fu.lead_id = cl.id
        LEFT JOIN users u ON u.id = cl.matchmaker_id
        WHERE {clause}
        GROUP BY cl.matchmaker_id, matchmaker_name
        ORDER BY total_count DESC, lead_count DESC"""), params)).mappings().all()
    return [MatchmakerFollowStat(
        matchmaker_id=int(r["matchmaker_id"]), matchmaker_name=str(r["matchmaker_name"] or ""),
        total_count=int(r["total_count"] or 0), month_count=int(r["month_count"] or 0),
        unfollowed_count=int(r["unfollowed_count"] or 0),
    ) for r in rows]


async def promoter_acquisition_stats(
    db: AsyncSession,
    *,
    promoter_id: int | None = None,
    scope: str | None = None,
    scope_params: dict[str, Any] | None = None,
) -> list[CustomerLeadPromoterStat]:
    """推广红娘获客统计（A13）：含男/女拆分。gender 列未迁移时男女相关列落 0。"""
    has_gender = "gender" in await _lead_columns(db)
    where = ["cl.promoter_id IS NOT NULL"]
    params: dict[str, Any] = {}
    if "deleted_at" in await _lead_columns(db):
        where.append("cl.deleted_at IS NULL")
    if scope:
        where.append(f"({scope})")
        params.update(scope_params or {})
    if promoter_id is not None:
        where.append("cl.promoter_id = :promoter_id")
        params["promoter_id"] = promoter_id
    clause = " AND ".join(where)
    if has_gender:
        # users 表也有 gender，JOIN 后必须带 cl. 前缀，否则 MySQL 1052 列歧义
        gender_cols = ("SUM(cl.gender = 1) male, SUM(cl.gender = 2) female, "
                       "SUM(cl.gender = 1 AND cl.audit_status = 'active' AND cl.status NOT IN ('LOST','CLOSED')) valid_male, "
                       "SUM(cl.gender = 2 AND cl.audit_status = 'active' AND cl.status NOT IN ('LOST','CLOSED')) valid_female")
    else:
        gender_cols = "0 male, 0 female, 0 valid_male, 0 valid_female"
    rows = (await db.execute(text(f"""SELECT cl.promoter_id,
        COALESCE(NULLIF(u.nickname, ''), u.phone, CONCAT('红娘', cl.promoter_id)) promoter_name,
        COUNT(*) total, {gender_cols},
        SUM(cl.audit_status = 'pending') pending,
        SUM(cl.status IN ('LOST','CLOSED')) invalid
        FROM customer_lead cl LEFT JOIN users u ON u.id = cl.promoter_id
        WHERE {clause}
        GROUP BY cl.promoter_id, promoter_name
        ORDER BY total DESC"""), params)).mappings().all()
    return [CustomerLeadPromoterStat(
        promoter_id=int(r["promoter_id"]), promoter_name=str(r["promoter_name"] or ""),
        total=int(r["total"] or 0), male=int(r["male"] or 0), female=int(r["female"] or 0),
        valid_male=int(r["valid_male"] or 0), valid_female=int(r["valid_female"] or 0),
        pending=int(r["pending"] or 0), invalid=int(r["invalid"] or 0),
    ) for r in rows]


async def abandon_lead(
    db: AsyncSession,
    account_id: int,
    lead_id: int,
    reason: str,
    *,
    scope: str | None = None,
    scope_params: dict[str, Any] | None = None,
) -> CustomerLeadAbandonment:
    lead = await get_lead(db, lead_id, scope=scope, scope_params=scope_params)
    if lead.status in ("CONVERTED", "CLOSED"):
        raise HTTPException(409, detail="已入库或已关闭的客源不能弃海")
    active = await db.execute(text("SELECT 1 FROM customer_lead_abandonment WHERE lead_id = :id AND restored_at IS NULL"), {"id": lead_id})
    if active.scalar():
        raise HTTPException(409, detail="该客源已在弃海池")
    result = await db.execute(text("""INSERT INTO customer_lead_abandonment (lead_id, reason, abandoned_by)
        VALUES (:lead_id, :reason, :account_id)"""), {"lead_id": lead_id, "reason": reason, "account_id": account_id})
    abandon_where = ["id = :id"]
    abandon_params: dict[str, Any] = {"id": lead_id}
    _merge_scope(abandon_where, abandon_params, scope, scope_params)
    abandon_assigned_sql = ", assigned_at = NULL" if "assigned_at" in await _lead_columns(db) else ""
    await db.execute(text(f"UPDATE customer_lead SET status = 'LOST', matchmaker_id = NULL{abandon_assigned_sql}, next_follow_at = NULL WHERE {' AND '.join(abandon_where)}"), abandon_params)
    await db.execute(text("INSERT INTO business_audit_log (actor_user_id, action, resource_type, resource_id) VALUES (:actor, 'customer_lead.abandon', 'customer_lead', :id)"), {"actor": account_id, "id": lead_id})
    await db.commit()
    row = (await db.execute(text("SELECT id, lead_id, reason, abandoned_by, abandoned_at, restored_by, restored_at, restore_reason FROM customer_lead_abandonment WHERE id = :id"), {"id": int(result.lastrowid)})).mappings().one()
    return CustomerLeadAbandonment(**dict(row))


async def restore_lead(
    db: AsyncSession,
    account_id: int,
    lead_id: int,
    reason: str,
    redaction_level: str = REDACTION_LEVEL_STANDARD,
    *,
    scope: str | None = None,
    scope_params: dict[str, Any] | None = None,
) -> CustomerLead:
    current = await get_lead(db, lead_id, scope=scope, scope_params=scope_params)
    active = (await db.execute(text("SELECT id FROM customer_lead_abandonment WHERE lead_id = :id AND restored_at IS NULL ORDER BY id DESC LIMIT 1"), {"id": lead_id})).scalar()
    if not active:
        raise HTTPException(409, detail="该客源不在弃海池")
    # 恢复后线索回到有效状态，先确认联系方式没有被其它有效线索占用。
    await ensure_contact_available(
        db,
        normalize_contact(current.phone),
        normalize_contact(current.wechat),
        exclude_lead_id=lead_id,
        detail="该客源的联系方式已被其他有效线索占用，无法恢复",
    )
    await db.execute(text("UPDATE customer_lead_abandonment SET restored_by = :account_id, restored_at = UTC_TIMESTAMP(), restore_reason = :reason WHERE id = :id"), {"account_id": account_id, "reason": reason, "id": active})
    restore_where = ["id = :id"]
    restore_params: dict[str, Any] = {"id": lead_id}
    _merge_scope(restore_where, restore_params, scope, scope_params)
    await db.execute(text(f"UPDATE customer_lead SET status = 'NEW', updated_at = UTC_TIMESTAMP() WHERE {' AND '.join(restore_where)}"), restore_params)
    await db.execute(text("INSERT INTO business_audit_log (actor_user_id, action, resource_type, resource_id) VALUES (:actor, 'customer_lead.restore', 'customer_lead', :id)"), {"actor": account_id, "id": lead_id})
    await db.commit()
    return await get_lead(db, lead_id, redaction_level, scope=scope, scope_params=scope_params)


async def delete_lead(
    db: AsyncSession,
    account_id: int,
    lead_id: int,
    *,
    scope: str | None = None,
    scope_params: dict[str, Any] | None = None,
) -> None:
    """软删除单条客源线索（A16）。

    仅对当前数据范围内、未删除的线索生效（越权/不存在/已删除均 404）；
    删除后生成列 active_phone/active_wechat 随 deleted_at 置 NULL，联系方式不再被占用。
    """
    if "deleted_at" not in await _lead_columns(db):
        raise HTTPException(503, detail="客源删除功能未就绪：请先执行软删除迁移")
    guard_where = ["id = :id", "deleted_at IS NULL"]
    guard_params: dict[str, Any] = {"id": lead_id}
    _merge_scope(guard_where, guard_params, scope, scope_params)
    if not await db.scalar(text(f"SELECT id FROM customer_lead WHERE {' AND '.join(guard_where)} FOR UPDATE"), guard_params):
        raise HTTPException(404, detail="客源线索不存在")
    await db.execute(text("""UPDATE customer_lead
        SET deleted_at = UTC_TIMESTAMP(), deleted_by = :account_id, updated_at = UTC_TIMESTAMP()
        WHERE id = :id AND deleted_at IS NULL"""), {"account_id": account_id, "id": lead_id})
    await db.execute(text("""INSERT INTO business_audit_log
        (actor_user_id, action, resource_type, resource_id, reason)
        VALUES (:actor, 'customer_lead.delete', 'customer_lead', :id, :reason)"""),
        {"actor": account_id, "id": lead_id, "reason": "后台软删除客源线索"})
    await db.commit()


async def batch_delete_leads(
    db: AsyncSession,
    account_id: int,
    lead_ids: list[int],
    *,
    scope: str | None = None,
    scope_params: dict[str, Any] | None = None,
) -> tuple[list[int], list[dict[str, str]]]:
    """批量软删除客源线索（A16），部分成功语义。

    单事务 + 一份汇总审计；越权/不存在/已删除的线索归入 failed，不整体报错。
    返回 (deleted_ids, [{id, error}, ...])。
    """
    ids = list(dict.fromkeys(int(item) for item in lead_ids))
    if "deleted_at" not in await _lead_columns(db):
        raise HTTPException(503, detail="客源删除功能未就绪：请先执行软删除迁移")
    if not ids:
        return [], []
    lock_params: dict[str, Any] = {}
    id_placeholders = _placeholders("bdel", [str(item) for item in ids], lock_params)
    lock_where = [f"id IN ({id_placeholders})", "deleted_at IS NULL"]
    _merge_scope(lock_where, lock_params, scope, scope_params)
    present = {
        int(row[0])
        for row in (await db.execute(text(
            f"SELECT id FROM customer_lead WHERE {' AND '.join(lock_where)} FOR UPDATE"), lock_params)).all()
    }
    deleted: list[int] = []
    failed: list[dict[str, str]] = []
    for lead_id in ids:
        if lead_id in present:
            deleted.append(lead_id)
        else:
            failed.append({"id": str(lead_id), "error": "客源不存在、已删除或越权"})
    if deleted:
        del_params: dict[str, Any] = {"account_id": account_id}
        del_placeholders = _placeholders("del", [str(item) for item in deleted], del_params)
        await db.execute(text(f"""UPDATE customer_lead
            SET deleted_at = UTC_TIMESTAMP(), deleted_by = :account_id, updated_at = UTC_TIMESTAMP()
            WHERE id IN ({del_placeholders}) AND deleted_at IS NULL"""), del_params)
        await db.execute(text("""INSERT INTO business_audit_log
            (actor_user_id, action, resource_type, resource_id, reason)
            VALUES (:actor, 'customer_lead.batch_delete', 'customer_lead', NULL, :reason)"""),
            {"actor": account_id, "reason": f"批量软删除 {len(deleted)} 条客源线索"})
    await db.commit()
    return deleted, failed


async def list_abandonments(
    db: AsyncSession,
    active_only: bool,
    *,
    scope: str | None = None,
    scope_params: dict[str, Any] | None = None,
) -> list[CustomerLeadAbandonment]:
    """弃海记录列表。

    弃海表自身不带归属列，归属在 ``customer_lead``；JOIN 后把作用域谓词限定在
    ``cl``（customer_lead）别名上，避免与弃海表同名列歧义（D-16）。
    """
    where = ["1=1"]
    params: dict[str, Any] = {}
    if active_only:
        where.append("ab.restored_at IS NULL")
    _merge_scope(where, params, scope, scope_params)
    rows = await db.execute(text(f"""SELECT ab.id, ab.lead_id, ab.reason, ab.abandoned_by, ab.abandoned_at,
        ab.restored_by, ab.restored_at, ab.restore_reason
        FROM customer_lead_abandonment ab JOIN customer_lead cl ON cl.id = ab.lead_id
        WHERE {' AND '.join(where)} ORDER BY ab.id DESC"""), params)
    return [CustomerLeadAbandonment(**dict(row)) for row in rows.mappings().all()]


async def _insert_leads(
    db: AsyncSession,
    account_id: int,
    rows: list[CustomerLeadImportRow],
    dup_mode: str = "skip",
    *,
    default_source: str = "批量导入",
    default_matchmaker_id: int | None = None,
    default_promoter_id: int | None = None,
    default_audit_status: str = "active",
    start_index: int = 2,
    assignment_scope: str | None = None,
    assignment_scope_params: dict[str, Any] | None = None,
) -> CustomerLeadImportSummary:
    """批量写入客源线索。留空/无法匹配的字段按 default_* 兜底；dup_mode=skip 时联系方式重复的行跳过。"""
    created = skipped = failed = 0
    errors: list[str] = []
    await _ensure_user_in_scope(db, default_matchmaker_id, assignment_scope, assignment_scope_params)
    await _ensure_user_in_scope(db, default_promoter_id, assignment_scope, assignment_scope_params)
    columns = await _lead_columns(db)
    for index, row in enumerate(rows, start=start_index):  # Excel 数据从第 2 行起
        if not row.phone and not row.wechat:
            failed += 1
            errors.append(f"第 {index} 行：phone 或 wechat 至少提供一个")
            continue
        try:
            if dup_mode == "skip":
                contact_conditions: list[str] = []
                dup_params: dict[str, Any] = {}
                if row.phone:
                    contact_conditions.append("phone = :phone")
                    dup_params["phone"] = row.phone
                if row.wechat:
                    contact_conditions.append("wechat = :wechat")
                    dup_params["wechat"] = row.wechat
                duplicate = await db.execute(
                    text("SELECT id FROM customer_lead WHERE status NOT IN ('LOST', 'CLOSED') AND ("
                         + " OR ".join(contact_conditions) + ") LIMIT 1"),
                    dup_params,
                )
                if duplicate.scalar():
                    skipped += 1
                    continue
            insert_columns = ["name", "phone", "wechat", "source", "intention_level", "remark",
                              "matchmaker_id", "promoter_id", "audit_status", "created_by"]
            insert_payload: dict[str, Any] = {
                "name": row.name, "phone": row.phone, "wechat": row.wechat,
                "source": (row.source or "").strip() or default_source,
                "intention_level": row.intention_level, "remark": row.remark,
                "matchmaker_id": default_matchmaker_id, "promoter_id": default_promoter_id,
                "audit_status": default_audit_status, "created_by": account_id,
            }
            # 导入即分派时补写分派时间，保证「分派日期」筛选覆盖导入产生的线索。
            if default_matchmaker_id and "assigned_at" in columns:
                insert_columns.append("assigned_at")
                insert_payload["assigned_at"] = datetime.utcnow()
            result = await db.execute(text(f"""INSERT INTO customer_lead
                ({', '.join(insert_columns)})
                VALUES (:{', '.join(insert_columns)})"""), insert_payload)
            lead_id = int(result.lastrowid)
            await db.execute(text("""INSERT INTO business_audit_log (actor_user_id, action, resource_type, resource_id)
                VALUES (:actor, 'customer_lead.import', 'customer_lead', :id)"""),
                {"actor": account_id, "id": lead_id},
            )
            created += 1
        except HTTPException:
            raise
        except Exception as exc:  # 单行失败不阻断整批
            failed += 1
            errors.append(f"第 {index} 行：{exc}")
    await db.commit()
    return CustomerLeadImportSummary(created=created, skipped=skipped, failed=failed, total=len(rows), errors=errors[:50])


async def batch_import_leads(
    db: AsyncSession,
    account_id: int,
    rows: list[CustomerLeadImportRow],
    dup_mode: str = "skip",
    *,
    assignment_scope: str | None = None,
    assignment_scope_params: dict[str, Any] | None = None,
) -> CustomerLeadBatchImportResult:
    """批量导入客源线索。dup_mode=skip 时联系方式重复的行跳过，append 时仍然导入。"""
    return await _insert_leads(db, account_id, rows, dup_mode, assignment_scope=assignment_scope, assignment_scope_params=assignment_scope_params)


# ─── Excel 模板下载与文件解析 ────────────────────────────────────────


def build_import_template() -> bytes:
    """生成「客源批量导入」Excel 模板（.xlsx）。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "客源导入"
    sheet.append(_TEMPLATE_HEADERS)
    sheet.append(_TEMPLATE_SAMPLE)
    for index, width in enumerate((18, 18, 18, 14, 12, 36), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    normalized = str(value).strip()
    return "" if normalized.lower() in {"nan", "none", "null"} else normalized


def _parse_intention(value: Any) -> int:
    raw = _cell_text(value)
    return _INTENTION_ALIASES.get(raw, 1)


def parse_import_file(content: bytes, filename: str) -> list[CustomerLeadImportRow]:
    """解析上传的 .xlsx 客源导入表，返回待写入的行（不做联系方式的即时校验）。"""
    lowered = (filename or "").lower()
    if lowered.endswith(".xls"):
        raise HTTPException(400, detail="暂不支持 .xls 旧格式，请用 Excel 另存为 .xlsx 后重新上传")
    if not lowered.endswith(".xlsx"):
        raise HTTPException(400, detail="仅支持 .xlsx 格式的 Excel 文件")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # 非 Excel / 文件损坏
        raise HTTPException(400, detail=f"Excel 解析失败：{exc}") from exc
    sheet = workbook.active
    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    headers = [_cell_text(cell) for cell in (header_row or ())]

    def find(*keywords: str) -> int | None:
        for index, header in enumerate(headers):
            if any(keyword in header for keyword in keywords):
                return index
        return None

    if find("姓名", "昵称", "name") is None or find("手机", "电话", "phone") is None:
        raise HTTPException(400, detail="模板表头缺少「姓名」或「手机号」列，请使用下载的模板填写")
    name_col, phone_col = find("姓名", "昵称", "name"), find("手机", "电话", "phone")
    wechat_col = find("微信", "wechat")
    source_col = find("来源", "source")
    intention_col = find("意向", "intention")
    remark_col = find("备注", "remark")

    rows: list[CustomerLeadImportRow] = []
    for raw in sheet.iter_rows(min_row=2, values_only=True):
        def pick(column: int | None) -> str:
            if column is None or column >= len(raw):
                return ""
            return _cell_text(raw[column])

        name = pick(name_col)
        phone = pick(phone_col)
        wechat = pick(wechat_col)
        if not name and not phone and not wechat:
            continue  # 整行空白，跳过
        if not name:
            raise HTTPException(400, detail="存在姓名为空的数据行，请补齐后再导入")
        rows.append(
            CustomerLeadImportRow(
                name=name,
                phone=phone or None,
                wechat=wechat or None,
                source=pick(source_col) or "批量导入",
                intention_level=_parse_intention(pick(intention_col)),
                remark=pick(remark_col) or None,
            )
        )
    if not rows:
        raise HTTPException(400, detail="文件中没有可导入的数据行")
    return rows


async def import_leads_file(
    db: AsyncSession,
    account_id: int,
    content: bytes,
    filename: str,
    *,
    audit_status: str = "active",
    default_source: str = "批量导入",
    default_matchmaker_id: int | None = None,
    default_promoter_id: int | None = None,
    dup_mode: str = "skip",
    assignment_scope: str | None = None,
    assignment_scope_params: dict[str, Any] | None = None,
) -> CustomerLeadImportSummary:
    """解析上传的 Excel 并按导入设置写入客源线索。"""
    rows = parse_import_file(content, filename)
    return await _insert_leads(
        db,
        account_id,
        rows,
        dup_mode,
        default_source=default_source or "批量导入",
        default_matchmaker_id=default_matchmaker_id,
        default_promoter_id=default_promoter_id,
        default_audit_status=audit_status,
        assignment_scope=assignment_scope,
        assignment_scope_params=assignment_scope_params,
    )


async def lead_options(
    db: AsyncSession,
    *,
    scope: str | None = None,
    scope_params: dict[str, Any] | None = None,
) -> CustomerLeadOptions:
    """客源线索页面下拉字典：来源、服务红娘、推广红娘。

    「来源」取自本账号可见的线索集合（带作用域，D-16）；红娘/推广为全局
    配置型字典，不随数据范围收敛。
    """
    source_where = ["source <> ''"]
    source_params: dict[str, Any] = {}
    _merge_scope(source_where, source_params, scope, scope_params)
    sources = [
        CustomerLeadOption(value=row["source"], label=row["source"])
        for row in (await db.execute(text(f"SELECT DISTINCT source FROM customer_lead WHERE {' AND '.join(source_where)} ORDER BY source"), source_params)).mappings().all()
    ]
    # A14：合并运营维护的来源字典（customer_lead_source）；字典表不存在（迁移未执行）时
    # 静默跳过，不阻断下拉。按 value 去重，字典优先展示，再补历史自由文本来源。
    source_dict_exists = (await db.execute(text("""SELECT 1 FROM information_schema.TABLES
        WHERE TABLE_SCHEMA = DATABASE() AND TABLE_NAME = 'customer_lead_source'"""))).scalar()
    dict_sources: list[CustomerLeadOption] = []
    if source_dict_exists:
        dict_sources = [
            CustomerLeadOption(value=row["name"], label=row["name"])
            for row in (await db.execute(text("SELECT name FROM customer_lead_source WHERE enabled = 1 ORDER BY sort, id"))).mappings().all()
        ]
    merged: dict[str, CustomerLeadOption] = {}
    for option in [*dict_sources, *sources]:
        merged.setdefault(option.value, option)
    sources = list(merged.values())
    matchmakers = [
        CustomerLeadOption(value=str(row["id"]), label=row["label"])
        for row in (await db.execute(text("""SELECT u.id, COALESCE(NULLIF(u.nickname, ''), u.phone, CONCAT('用户', u.id)) AS label
            FROM user_matchmaker_apply a JOIN users u ON u.id = a.user_id
            JOIN user_role r ON r.user_id = a.user_id AND r.role_code = 'service_matchmaker' AND r.status = 1
            WHERE a.application_type = 'service_matchmaker' AND a.status = 1 ORDER BY u.id"""))).mappings().all()
    ]
    promoters = [
        CustomerLeadOption(value=str(row["id"]), label=row["label"])
        for row in (await db.execute(text("""SELECT u.id, COALESCE(NULLIF(u.nickname, ''), u.phone, CONCAT('用户', u.id)) AS label
            FROM user_matchmaker_apply a JOIN users u ON u.id = a.user_id
            WHERE a.application_type = 'promoter' AND a.status = 1 ORDER BY u.id"""))).mappings().all()
    ]
    return CustomerLeadOptions(sources=sources, matchmakers=matchmakers, promoters=promoters)
